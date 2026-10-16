"""agentkit.textfassung mit den Bibliotheken dieser .venv (ohne pypdf: PDFs liest Claude direkt)."""

import tempfile
import unittest
from pathlib import Path

import docx
import openpyxl
import pptx

from agentkit.textfassung import als_text


class TestTextfassung(unittest.TestCase):
    def test_office_formate(self):
        tmp = Path(tempfile.mkdtemp())
        d = docx.Document()
        d.add_paragraph("Antrag")
        d.add_table(rows=1, cols=2).rows[0].cells[0].text = "Betrag"
        d.save(tmp / "a.docx")
        self.assertEqual(als_text(tmp / "a.docx"), "Antrag\nBetrag | ")

        p = pptx.Presentation()
        p.slides.add_slide(p.slide_layouts[5]).shapes.title.text = "Ziele"
        p.save(tmp / "b.pptx")
        self.assertEqual(als_text(tmp / "b.pptx"), "--- Folie 1\nZiele")

        w = openpyxl.Workbook()
        w.active.title = "Kosten"
        w.active.append(["Holz", 120])
        w.save(tmp / "c.xlsx")
        self.assertEqual(als_text(tmp / "c.xlsx"), "--- Blatt Kosten\nHolz | 120")

        self.assertIsNone(als_text(tmp / "bild.png"))


if __name__ == "__main__":
    unittest.main()
