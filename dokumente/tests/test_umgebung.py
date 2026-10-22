"""Prüft, dass die .venv vollständig ist: agentkit, Renderer und die Office-Bibliotheken."""

import tempfile
import unittest
from pathlib import Path

import docx
import openpyxl
import pptx
import pypdfium2  # noqa: F401 – Vorschau des Renderers
from PIL import Image  # noqa: F401

from agentkit import auftraege  # noqa: F401 – Einbindung per .pth
import renderer  # noqa: F401 – ebenfalls per .pth


class TestUmgebung(unittest.TestCase):
    def test_office_dateien_schreiben_und_lesen(self):
        tmp = Path(tempfile.mkdtemp())
        d = docx.Document()
        d.add_heading("Titel", level=1)
        d.save(tmp / "a.docx")
        self.assertEqual(docx.Document(tmp / "a.docx").paragraphs[0].text, "Titel")

        p = pptx.Presentation()
        p.slides.add_slide(p.slide_layouts[0]).shapes.title.text = "Folie"
        p.save(tmp / "a.pptx")
        self.assertEqual(pptx.Presentation(tmp / "a.pptx").slides[0].shapes.title.text, "Folie")

        w = openpyxl.Workbook()
        w.active["A1"], w.active["A2"] = 2, "=A1*3"
        w.save(tmp / "a.xlsx")
        self.assertEqual(openpyxl.load_workbook(tmp / "a.xlsx").active["A2"].value, "=A1*3")


if __name__ == "__main__":
    unittest.main()
