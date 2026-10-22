"""Zusammenfassung einer Vorlesung als docx bauen, aus einer geprüften Beschreibung (pruefung.lesen).

Aufbau: Kopf (Kurs, Titel, Vorlesung vom, Stand, erstellt von, Legende der Kennzeichen), nummerierte
Überschriften, Absätze, Listen, Tabellen, Hinweiskästen, Rechenbeispiele; am Ende Selbsttest, offene Fragen
und auf eigener Seite die Lösungen. Jede Aussage kann Quellen und ein Kennzeichen tragen, beides steht klein
am Ende des Absatzes. Fußzeile: Kurs, Titel, Seite x von y (außer die Vorlage hat eine eigene).
Seite, Schrift und XML-Hilfen kommen aus arbeitsblatt.py, dessen Verhalten hier nicht geändert wird.
"""

import datetime
from pathlib import Path

from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from renderer import arbeitsblatt, design
from renderer.arbeitsblatt import (_breiten, _xml_absatz_farbe, _xml_absatz_rahmen, _xml_rahmen,
                                   _xml_zeile_wiederholen, _xml_zelle_farbe)

UEBERSCHRIFT_PT = {1: 14, 2: 12}
TITEL_PT = 20
KLEIN_PT = 9  # Kopfzeile, Quellen, Kennzeichen, Fußzeile
ABSTAND_NACH_PT = 6
EINZUG_CM = 0.63
BALKEN = ("24", "6")  # linker Balken der Kästen: Stärke in Achtelpunkten, Abstand in Punkten
KENNZEICHEN = {"unsicher": "unsicher", "muendlich": "nur mündlich"}
LEGENDE = {"unsicher": "Aussage im Transkript nicht eindeutig, bitte prüfen",
           "muendlich": "nur in der Vorlesung gesagt, nicht auf den Folien"}
HINWEIS_TITEL = {"pruefungsrelevant": "Prüfungsrelevant", "organisatorisch": "Organisatorisches"}
QUELLE_VOR = {"transkript": "Transkript S. ", "folie": "Folie ", "norm": ""}


def heute() -> str:
    """Stand der Zusammenfassung (eigene Funktion, damit Tests sie festsetzen können)."""
    return datetime.date.today().strftime("%d.%m.%Y")


def bauen(beschreibung: dict, ziel: Path, profil: dict | None = None, vorlage: Path | None = None) -> Path:
    """Baut die Zusammenfassung nach `ziel` (.docx). `beschreibung` aus pruefung.lesen(), `profil` siehe design.py,
    `vorlage` eine .docx (Kopfzeile mit Logo, Ränder; Schriften und Farben kommen weiter aus dem Profil)."""
    b = _Bau(design.profil_lesen(profil), vorlage)
    b.kopf(beschreibung)
    for block in beschreibung["bloecke"]:
        getattr(b, block["typ"])(block)
    b.schluss(beschreibung["selbsttest"], beschreibung["offene_fragen"])
    b.fusszeile(beschreibung)
    b.d.save(ziel)
    return ziel


class _Bau(arbeitsblatt._Bau):
    """Wie das Arbeitsblatt, mit eigenen Überschriften, Kästen und Belegen. Zählt die Überschriften."""

    def __init__(self, profil: dict, vorlage: Path | None):
        super().__init__(profil, vorlage)
        self.nummer = [0, 0]
        self.farben = profil["farben"]

    def _stile(self):
        super()._stile()
        self.d.styles["Normal"].paragraph_format.space_after = Pt(ABSTAND_NACH_PT)
        for ebene, groesse in UEBERSCHRIFT_PT.items():
            stil = self.d.styles[f"Heading {ebene}"]
            stil.font.size = Pt(groesse)
            stil.font.color.rgb = RGBColor.from_string(self.profil["farben"]["ueberschrift"])
            stil.paragraph_format.space_before = Pt(18 if ebene == 1 else 12)
            stil.paragraph_format.keep_with_next = True

    # Kopf und Fußzeile

    def kopf(self, b: dict):
        p = self.profil
        self.d.core_properties.title = b["titel"]
        self.d.core_properties.author = p["erstellt_von"]
        self._zeile(b["kurs"], KLEIN_PT + 2, fett=True, farbe=self.farben["ueberschrift"], nach=2)
        self._zeile(b["titel"], TITEL_PT, fett=True, nach=2 if b["untertitel"] else 6)
        if b["untertitel"]:
            self._zeile(b["untertitel"], UEBERSCHRIFT_PT[2], nach=6)
        angaben = [f"Vorlesung vom {b['datum']}", f"Stand {heute()}"]
        if p["erstellt_von"]:
            angaben.append(f"erstellt von {p['erstellt_von']}")
        self._zeile(" · ".join(angaben), KLEIN_PT + 1, nach=2)
        benutzt = [k for k in KENNZEICHEN if k in _kennzeichen(b)]
        if benutzt:
            zeile = self.d.add_paragraph()
            zeile.paragraph_format.space_after = Pt(2)
            for i, k in enumerate(benutzt):
                self._kennzeichen(zeile, k, davor="" if i == 0 else "   ")
                self._klein(zeile, f" {LEGENDE[k]}")
        trenner = self.d.add_paragraph()  # leer, nur die Linie darunter
        trenner.paragraph_format.space_after = Pt(6)
        _xml_absatz_rahmen(trenner, "single", "6", self.farben["ueberschrift"], ("bottom",), abstand="1")

    def fusszeile(self, b: dict):
        fuss = self.d.sections[0].footer
        if self.vorlage and _hat_inhalt(fuss):
            return  # Fußzeile der Vorlage behalten
        p = fuss.paragraphs[0]
        tabs = p.paragraph_format.tab_stops
        for stopp in self.d.styles["Footer"].paragraph_format.tab_stops:  # Mitte und rechts des Stils löschen
            tabs.add_tab_stop(stopp.position, WD_TAB_ALIGNMENT.CLEAR)
        tabs.add_tab_stop(Cm(self.textbreite_cm - 0.01), WD_TAB_ALIGNMENT.RIGHT)  # genau am Rand übergeht LibreOffice
        self._klein(p, f"{b['kurs']} · {b['titel']}\tSeite ")
        self._feld(p, "PAGE")
        self._klein(p, " von ")
        self._feld(p, "NUMPAGES")

    # Blöcke (Name = typ in der Beschreibung)

    def ueberschrift(self, block: dict):
        ebene = block["ebene"]
        self.nummer[ebene - 1] += 1
        if ebene == 1:
            self.nummer[1] = 0
        nummer = ".".join(str(n) for n in self.nummer[:ebene])
        self.d.add_heading(f"{nummer}  {block['text']}", level=ebene)

    def absatz(self, block: dict):
        self._beleg(self.d.add_paragraph(block["text"]), block)

    def liste(self, block: dict):
        name = "List Number" if block["nummeriert"] else "List Bullet"
        stil = self._stil(name)
        ohne_stil = stil is None or stil.element.pPr is None or stil.element.pPr.numPr is None
        absaetze = []
        for i, punkt in enumerate(block["punkte"], 1):
            if ohne_stil:  # Vorlage ohne Listenstil: Zeichen von Hand, mit hängendem Einzug
                p = self.d.add_paragraph(f"{i}.\t" if block["nummeriert"] else "•\t")
                p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(EINZUG_CM), Cm(-EINZUG_CM)
                p.paragraph_format.tab_stops.add_tab_stop(Cm(EINZUG_CM))
                p.add_run(punkt["text"])
            else:
                p = self.d.add_paragraph(punkt["text"], style=name)
            p.paragraph_format.space_after = Pt(2)
            self._beleg(p, punkt)
            absaetze.append(p)
        absaetze[-1].paragraph_format.space_after = Pt(ABSTAND_NACH_PT)
        if block["nummeriert"] and not ohne_stil:
            arbeitsblatt._xml_nummerierung_neu(self.d, name, absaetze)

    def tabelle(self, block: dict):
        if block["titel"]:
            titel = self.d.add_paragraph()
            titel.add_run(block["titel"]).bold = True
            titel.paragraph_format.space_after = Pt(3)
            titel.paragraph_format.keep_with_next = True
        zeilen = ([block["kopf"]] if block["kopf"] else []) + block["zeilen"]
        spalten = len(zeilen[0])
        t = self.d.add_table(rows=len(zeilen), cols=spalten)
        _xml_rahmen(t)
        _breiten(t, [self.textbreite_cm / spalten] * spalten)
        for nr, (reihe, werte) in enumerate(zip(t.rows, zeilen, strict=True)):
            ist_kopf = block["kopf"] is not None and nr == 0
            reihe.height, reihe.height_rule = Cm(0.6), WD_ROW_HEIGHT_RULE.AT_LEAST
            for spalte, (zelle, wert) in enumerate(zip(reihe.cells, werte, strict=True)):
                p = zelle.paragraphs[0]
                p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(2)
                run = p.add_run(wert)
                run.bold = ist_kopf or (block["erste_spalte_fett"] and spalte == 0)
                if ist_kopf:
                    _xml_zelle_farbe(zelle, self.farben["tabellenkopf"])
        if block["kopf"]:
            _xml_zeile_wiederholen(t.rows[0])
        unter = self.d.add_paragraph()  # trennt auch zwei Tabellen (siehe arbeitsblatt._abstand)
        self._beleg(unter, block)

    def hinweis(self, block: dict):
        titel = block["titel"] or HINWEIS_TITEL[block["art"]]
        absaetze = self._kasten(titel, block["text"].split("\n"), self.farben[block["art"]])
        self._beleg(absaetze[-1], block)
        self._nach_kasten()

    def rechenbeispiel(self, block: dict):
        titel = "Rechenbeispiel" + (f": {block['titel']}" if block["titel"] else "")
        absaetze = self._kasten(titel, block["aufgabe"].split("\n"), self.farben["rechenbeispiel"])
        if block["gegeben"]:
            absaetze.append(self._im_kasten("Gegeben:", fett=True))
            absaetze += [self._im_kasten(f"•\t{g}", einzug=EINZUG_CM) for g in block["gegeben"]]
        for nr, schritt in enumerate(block["schritte"], 1):
            p = self._im_kasten("")
            p.add_run(f"Schritt {nr}: ").bold = True
            p.add_run(schritt["text"])
            absaetze.append(p)
            if schritt["rechnung"]:
                absaetze.append(self._im_kasten(f"\t{schritt['rechnung']}", einzug=EINZUG_CM * 2))
        ergebnis = self._im_kasten("")
        ergebnis.add_run("Ergebnis: ").bold = True
        ergebnis.add_run(block["ergebnis"]).bold = True
        absaetze.append(ergebnis)
        for p in absaetze:
            self._kasten_form(p, self.farben["rechenbeispiel"])
        self._beleg(ergebnis, block)
        self._nach_kasten()

    def schluss(self, selbsttest: list, offene: list):
        if selbsttest:
            self._abschnitt("Selbsttest")
            for nr, frage in enumerate(selbsttest, 1):
                self._nummeriert(nr, frage["frage"])
        if offene:
            self._abschnitt("Offene Fragen")
            self.liste({"nummeriert": False, "punkte": offene})
        if selbsttest:
            self._abschnitt("Lösungen zum Selbsttest", neue_seite=True)
            for nr, frage in enumerate(selbsttest, 1):
                self._beleg(self._nummeriert(nr, frage["loesung"]), frage)

    # Hilfen

    def _zeile(self, text: str, groesse: float, fett: bool = False, farbe: str | None = None, nach: float = 0):
        p = self.d.add_paragraph()
        run = p.add_run(text)
        run.font.size, run.bold = Pt(groesse), fett
        if farbe:
            run.font.color.rgb = RGBColor.from_string(farbe)
        p.paragraph_format.space_after = Pt(nach)
        return p

    def _abschnitt(self, text: str, neue_seite: bool = False):
        """Überschrift ohne Nummer für Selbsttest, offene Fragen, Lösungen."""
        p = self.d.add_heading(text, level=1)
        if neue_seite:  # False schriebe ein Element mit val="0"
            p.paragraph_format.page_break_before = True

    def _nummeriert(self, nr: int, text: str):
        p = self.d.add_paragraph()
        p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(EINZUG_CM), Cm(-EINZUG_CM)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(EINZUG_CM))
        p.add_run(f"{nr}.").bold = True
        p.add_run(f"\t{text}")
        return p

    def _kasten(self, titel: str, zeilen: list[str], farbe: str) -> list:
        """Titel fett und je Zeile ein Absatz, alle mit Fläche und Balken links (Word fasst sie zusammen)."""
        absaetze = [self._im_kasten(titel, fett=True)] + [self._im_kasten(z) for z in zeilen]
        for p in absaetze:
            self._kasten_form(p, farbe)
        return absaetze

    def _im_kasten(self, text: str, fett: bool = False, einzug: float = 0):
        p = self.d.add_paragraph()
        if text:
            p.add_run(text).bold = fett
        f = p.paragraph_format
        f.space_after = Pt(2)
        f.keep_with_next = True
        if einzug:  # hängend: die erste Zeile beginnt bei 0, sonst rückt der Balken links mit ein
            f.left_indent, f.first_line_indent = Cm(einzug), Cm(-einzug)
            f.tab_stops.add_tab_stop(Cm(einzug))
        return p

    def _kasten_form(self, p, farbe: str):
        if p._p.pPr is not None and p._p.pPr.find(qn("w:pBdr")) is not None:
            return
        _xml_absatz_rahmen(p, "single", BALKEN[0], self.farben["ueberschrift"], ("left",), abstand=BALKEN[1])
        _xml_absatz_farbe(p, farbe)

    def _nach_kasten(self):
        """Leerer Absatz ohne Rahmen: trennt den Kasten vom nächsten (gleiche Rahmen verschmelzen)."""
        p = self.d.add_paragraph()
        p.paragraph_format.space_after = Pt(0)

    def _beleg(self, p, d: dict):
        """Kennzeichen und Quellen klein ans Ende des Absatzes."""
        if d.get("kennzeichnung"):  # Lösungen im Selbsttest haben keins
            self._kennzeichen(p, d["kennzeichnung"], davor=" " if p.text else "")
        if d["quellen"]:
            text = "; ".join(QUELLE_VOR[q["art"]] + q["stelle"] for q in d["quellen"])
            self._klein(p, f"{' ' if p.text else ''}[{text}]")

    def _feld(self, p, befehl: str):
        """Feld wie Seitenzahl (PAGE) oder Seitenanzahl (NUMPAGES), je Teil ein Run wie bei Word selbst."""
        for teil in ("begin", befehl, "separate", "1", "end"):
            run = self._klein(p, teil if teil == "1" else "")
            if teil in ("begin", "separate", "end"):
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), teil)
                run._r.append(el)
            elif teil == befehl:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = f" {befehl} "
                run._r.append(el)

    def _kennzeichen(self, p, art: str, davor: str):
        run = self._klein(p, f"{davor}[{KENNZEICHEN[art]}]")
        run.bold = True
        return run

    def _klein(self, p, text: str):
        run = p.add_run(text)
        run.font.size = Pt(KLEIN_PT)
        run.font.color.rgb = RGBColor.from_string(self.farben["kennzeichnung"])
        return run


def _kennzeichen(b: dict) -> set:
    """Alle Kennzeichen der Beschreibung (für die Legende im Kopf)."""
    gefunden = set()
    for block in b["bloecke"]:
        for d in [block] + block.get("punkte", []):
            if d.get("kennzeichnung"):
                gefunden.add(d["kennzeichnung"])
    gefunden.update(f["kennzeichnung"] for f in b["offene_fragen"] if f["kennzeichnung"])
    return gefunden


def _hat_inhalt(teil) -> bool:
    """Kopf- oder Fußzeile mit Text oder Bild."""
    return any(p.text.strip() for p in teil.paragraphs) or bool(teil._element.findall(".//" + qn("w:drawing")))

