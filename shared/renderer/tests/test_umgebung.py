"""Prüft, dass die .venv vollständig ist: renderer und agentkit per .pth, docx und pptx."""

import tempfile
import unittest
from pathlib import Path

import docx
import pptx

import renderer  # noqa: F401 – Einbindung per .pth
from agentkit import textfassung  # noqa: F401 – für die Textfassung zur Selbstkontrolle


class TestUmgebung(unittest.TestCase):
    def test_office_dateien_schreiben_und_lesen(self):
        tmp = Path(tempfile.mkdtemp())
        d = docx.Document()
        d.add_heading("Titel", level=1)
        d.save(tmp / "a.docx")
        self.assertEqual(docx.Document(tmp / "a.docx").paragraphs[0].text, "Titel")

        p = pptx.Presentation()
        p.slides.add_slide(p.slide_layouts[0]).shapes.title.text = "Folie"
        p.save(tmp / "a.pptx")
        self.assertEqual(pptx.Presentation(tmp / "a.pptx").slides[0].shapes.title.text, "Folie")


if __name__ == "__main__":
    unittest.main()
