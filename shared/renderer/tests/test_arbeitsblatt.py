"""Arbeitsblatt als docx: bauen, wieder öffnen, Inhalt und Aufbau prüfen. Das Aussehen prüft die Vorschau."""

import itertools
import tempfile
import unittest
from pathlib import Path

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

from renderer import arbeitsblatt, design
from renderer.pruefung import lesen

PROFIL = {"fach": "Ethik", "lehrkraft": "Frau Muster", "farben": {"thema": "FFD873", "merksatz": "F2CEED",
                                                                    "material": "D9F2D0"}}

BLATT = {
    "art": "arbeitsblatt",
    "titel": "Muss ich helfen?",
    "nummer": "5",
    "bloecke": [
        {"typ": "merksatz", "titel": "Merksatz:", "text": "Ist die Not groß, helfe ich."},
        {"typ": "anweisung", "text": "Lies und bearbeite!"},
        {"typ": "material", "titel": "Beispiel", "text": "Erste Zeile\nzweite Zeile"},
        {"typ": "aufgaben", "aufgaben": [
            {"text": "Unterstreiche!", "hilfe": "Zeilen 1-2"},
            {"text": "Kreuze an!", "antwort": {"art": "ankreuzen", "optionen": ["A", "B"]}},
            {"text": "Begründe!", "antwort": {"art": "satzanfaenge", "saetze": [
                {"anfang": "Ich finde", "linien": 2}, {"anfang": "Zum Beispiel", "linien": 1}]}, "hilfe": "BBB"},
            {"text": "Erkläre.", "antwort": {"art": "linien", "anzahl": 3}},
            {"text": "Zeichne.", "antwort": {"art": "kasten", "hoehe_cm": 6.5}},
        ]},
        {"typ": "ueberschrift", "text": "Vertiefung", "ebene": 2},
        {"typ": "liste", "punkte": ["eins", "zwei"], "nummeriert": True},
        {"typ": "liste", "punkte": ["noch eins"], "nummeriert": True},
        {"typ": "aufgaben", "titel": "", "aufgaben": [{"text": "Weiter."}]},
        {"typ": "neue_seite", "querformat": True},
        {"typ": "tabelle", "kopf": ["", "Singer"], "zeilen": [["Behauptung", ""], ["Beispiel", ""]],
         "zeilenhoehe_cm": 4, "erste_spalte_fett": True},
    ],
}


def bauen(beschreibung=BLATT, profil=PROFIL) -> docx.document.Document:
    ziel = Path(tempfile.mkdtemp()) / "blatt.docx"
    arbeitsblatt.bauen(lesen(beschreibung), ziel, profil)
    return docx.Document(ziel)


def cm(laenge) -> float:
    """Word speichert in Zwanzigstel-Punkt: auf 0,01 cm vergleichen."""
    return round(laenge.cm, 2)


def num_id(p):
    return p._p.pPr.numPr.numId.val


def finden(d, anfang: str):
    return [p for p in d.paragraphs if p.text.startswith(anfang)][0]


def ppr(p, name: str):
    return p._p.pPr.find(qn(name)) if p._p.pPr is not None else None


class TestArbeitsblatt(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.d = bauen()

    def test_seiten(self):
        hoch, quer = self.d.sections
        self.assertEqual((cm(hoch.page_width), cm(hoch.page_height)), (21, 29.7))
        self.assertEqual(cm(hoch.left_margin), 2)
        self.assertEqual(quer.orientation, WD_ORIENT.LANDSCAPE)
        self.assertEqual((cm(quer.page_width), cm(quer.page_height)), (29.7, 21))

    def test_eigenschaften_und_schrift(self):
        self.assertEqual(self.d.core_properties.title, "Muss ich helfen?")
        self.assertEqual(self.d.core_properties.author, "")
        normal = self.d.styles["Normal"]
        self.assertEqual((normal.font.name, normal.font.size.pt), ("Calibri", 11))
        rfonts = normal.element.rPr.rFonts
        self.assertIsNone(rfonts.get(qn("w:asciiTheme")))  # Theme-Schrift hätte Vorrang
        self.assertEqual(normal.element.rPr.find(qn("w:lang")).get(qn("w:val")), "de-DE")
        h2 = self.d.styles["Heading 2"]
        self.assertEqual((h2.font.name, h2.font.size.pt, str(h2.font.color.rgb)), ("Calibri", 12, "000000"))

    def test_kopf(self):
        t = self.d.tables[0]
        texte = [[z.text for z in r.cells] for r in t.rows]
        self.assertEqual(texte, [["Ethik", "Muss ich helfen?", "Datum:"],
                                 ["Ethik", "Muss ich helfen?", "Name:"],
                                 ["Frau Muster", "Arbeitsblatt 5", "Name:"]])
        shd = t.cell(0, 1)._tc.tcPr.find(qn("w:shd"))
        self.assertEqual(shd.get(qn("w:fill")), "FFD873")

    def test_kopf_ohne_profil(self):
        d = bauen({k: v for k, v in BLATT.items() if k != "nummer"}, profil=None)
        t = d.tables[0]
        self.assertEqual([t.cell(2, 0).text, t.cell(2, 1).text], ["", "Arbeitsblatt"])
        self.assertEqual(t.cell(0, 1)._tc.tcPr.find(qn("w:shd")).get(qn("w:fill")),
                         design.STANDARD["farben"]["thema"])

    def folgende(self, anfang: str, anzahl: int):
        texte = [p.text for p in self.d.paragraphs]
        i = texte.index(anfang)
        return self.d.paragraphs[i:i + anzahl]

    def test_merksatz_im_kasten(self):
        titel, text = self.folgende("Merksatz:", 2)
        self.assertEqual(text.text, "Ist die Not groß, helfe ich.")
        self.assertTrue(titel.runs[0].bold)
        for p in (titel, text):  # gleicher Rahmen: ein Kasten
            rahmen = ppr(p, "w:pBdr")
            self.assertEqual({k.tag.split("}")[1] for k in rahmen}, {"top", "left", "bottom", "right"})
            self.assertEqual(rahmen.find(qn("w:top")).get(qn("w:val")), "thickThinSmallGap")
            self.assertEqual(ppr(p, "w:shd").get(qn("w:fill")), "F2CEED")

    def test_anweisung_und_material(self):
        self.assertTrue(finden(self.d, "Lies und").runs[0].bold)
        absaetze = self.folgende("Beispiel", 3)
        self.assertEqual([p.text for p in absaetze], ["Beispiel", "Erste Zeile", "zweite Zeile"])  # kein \n im Absatz
        for p in absaetze:
            self.assertEqual(ppr(p, "w:shd").get(qn("w:fill")), "D9F2D0")
            self.assertEqual(p.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
            self.assertIsNone(ppr(p, "w:pBdr"))

    def test_aufgaben_durchnummeriert(self):
        texte = [p.text for p in self.d.paragraphs if p.text[:2] in ("1.", "2.", "3.", "4.", "5.", "6.")]
        self.assertEqual(texte, ["1.\tUnterstreiche!", "2.\tKreuze an!", "3.\tBegründe!", "4.\tErkläre.",
                                 "5.\tZeichne.", "6.\tWeiter."])
        self.assertEqual([p.text for p in self.d.paragraphs].count("Aufgaben:"), 1)  # zweite Gruppe ohne Titel
        p = finden(self.d, "1.\t")
        self.assertTrue(p.runs[0].bold)
        self.assertEqual((cm(p.paragraph_format.left_indent), cm(p.paragraph_format.first_line_indent)), (1.3, -0.7))
        self.assertFalse(p.paragraph_format.keep_with_next)  # ohne Antwortplatz
        self.assertTrue(finden(self.d, "2.\t").paragraph_format.keep_with_next)

    def test_ankreuzen(self):
        p = self.d.paragraphs[[x.text for x in self.d.paragraphs].index("2.\tKreuze an!") + 1]
        symbole = p._p.findall(".//" + qn("w:sym"))
        self.assertEqual([(s.get(qn("w:font")), s.get(qn("w:char"))) for s in symbole], [("Wingdings", "F070")] * 2)
        self.assertIn(" A", p.text)
        self.assertIn(" B", p.text)

    def test_satzanfaenge(self):
        texte = [p.text for p in self.d.paragraphs]
        i = texte.index("Ich finde")
        self.assertEqual(texte[i + 3], "Zum Beispiel")
        for p in self.d.paragraphs[i + 1:i + 3]:
            self.assertIsNotNone(ppr(p, "w:pBdr"))

    def test_hilfen_klein_nach_der_gruppe(self):
        texte = [p.text for p in self.d.paragraphs]
        i = texte.index("Aufgabe 1:\tZeilen 1-2")
        self.assertEqual(texte[i + 1], "Aufgabe 3:\tBBB")
        self.assertGreater(i, texte.index("5.\tZeichne."))
        p = self.d.paragraphs[i]
        self.assertEqual({r.font.size.pt for r in p.runs}, {9})
        self.assertTrue(p.runs[0].bold)

    def test_antwortlinien(self):
        i = [p.text for p in self.d.paragraphs].index("4.\tErkläre.")
        linien = self.d.paragraphs[i + 1:i + 4]
        for p in linien:
            rahmen = ppr(p, "w:pBdr")
            self.assertIsNotNone(rahmen.find(qn("w:bottom")))
            self.assertIsNotNone(rahmen.find(qn("w:between")))  # sonst nur eine Linie unter der Gruppe
            self.assertEqual(cm(p.paragraph_format.line_spacing), 0.8)
            self.assertEqual(p.paragraph_format.line_spacing_rule, WD_LINE_SPACING.EXACTLY)
        self.assertEqual([p.paragraph_format.keep_with_next for p in linien], [True, True, False])

    def test_kasten(self):
        kasten = self.d.tables[1]
        self.assertEqual((len(kasten.rows), len(kasten.columns)), (1, 1))
        self.assertEqual(cm(kasten.rows[0].height), 6.5)
        self.assertEqual(kasten.rows[0].height_rule, WD_ROW_HEIGHT_RULE.EXACTLY)

    def test_listen_beginnen_neu(self):
        nummeriert = [p for p in self.d.paragraphs if p.style.name == "List Number"]
        self.assertEqual([p.text for p in nummeriert], ["eins", "zwei", "noch eins"])
        self.assertEqual(num_id(nummeriert[0]), num_id(nummeriert[1]))
        self.assertNotEqual(num_id(nummeriert[1]), num_id(nummeriert[2]))

    def test_ausfuelltabelle_im_querformat(self):
        t = self.d.tables[2]
        self.assertEqual([[z.text for z in r.cells] for r in t.rows], [["", "Singer"], ["Behauptung", ""],
                                                                        ["Beispiel", ""]])
        self.assertIsNone(t.rows[0].height)
        self.assertEqual(cm(t.rows[1].height), 4)
        self.assertTrue(t.rows[1].cells[0].paragraphs[0].runs[0].bold)
        self.assertFalse(t.rows[1].cells[1].paragraphs[0].runs[0].bold)
        self.assertIsNotNone(t.rows[0]._tr.trPr.find(qn("w:tblHeader")))
        self.assertEqual(cm(t.rows[1].cells[0].width), round(25.7 / 2, 2))  # Textbreite quer

    def test_tabellen_nie_direkt_hintereinander(self):
        koerper = list(self.d.element.body)
        for vorher, nachher in itertools.pairwise(koerper):
            self.assertFalse(vorher.tag == nachher.tag == qn("w:tbl"))

    def test_xml_reihenfolge_nach_schema(self):
        """Word lehnt Dateien mit falscher Reihenfolge ab: pBdr vor shd vor spacing, tblBorders vor tblLook."""
        reihenfolge = [qn(n) for n in ("w:keepNext", "w:pBdr", "w:shd", "w:tabs", "w:spacing", "w:ind")]
        for p in self.d.paragraphs:
            if p._p.pPr is None:
                continue
            namen = [k.tag for k in p._p.pPr if k.tag in reihenfolge]
            self.assertEqual(namen, sorted(namen, key=reihenfolge.index), p.text[:20])
        for t in self.d.tables:
            namen = [k.tag for k in t._tbl.tblPr]
            self.assertLess(namen.index(qn("w:tblBorders")), namen.index(qn("w:tblLook")))


class TestProfil(unittest.TestCase):
    def test_standard_und_ueberschreiben(self):
        p = design.profil_lesen({"fach": "Ethik", "farben": {"thema": "#ffd873"}})
        self.assertEqual(p["fach"], "Ethik")
        self.assertEqual(p["farben"]["thema"], "FFD873")
        self.assertEqual(p["farben"]["merksatz"], design.STANDARD["farben"]["merksatz"])
        self.assertEqual(design.profil_lesen(None), design.STANDARD)

    def test_fehler(self):
        with self.assertRaises(design.ProfilUngueltig) as ctx:
            design.profil_lesen({"farbe": 1, "farben": {"thema": "gelb", "rand": "000000"}, "linienabstand_cm": 3})
        meldung = str(ctx.exception)
        self.assertIn("farbe: unbekannt", meldung)
        self.assertIn('farben.thema: muss eine Farbe wie "FFD873" sein', meldung)
        self.assertIn("farben.rand: unbekannt", meldung)
        self.assertIn("linienabstand_cm: muss eine Zahl von 0.5 bis 1.5 sein", meldung)


if __name__ == "__main__":
    unittest.main()
