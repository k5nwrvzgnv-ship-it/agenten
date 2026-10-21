"""Vorlagen: Umwandlung .dotx/.potx, Word- und PowerPoint-Vorlage als Grundlage, Rückfall bei fehlenden Stilen."""

import tempfile
import unittest
import zipfile
from pathlib import Path

import docx
import pptx
from docx.shared import Cm
from pptx.util import Emu

from renderer import arbeitsblatt, folien, vorlage_umwandeln, vorlagen
from renderer.pruefung import lesen

BLATT = lesen({"art": "arbeitsblatt", "titel": "Thema", "bloecke": [
    {"typ": "ueberschrift", "text": "Teil 1"},
    {"typ": "liste", "punkte": ["a", "b"], "nummeriert": True},
    {"typ": "liste", "punkte": ["c"]},
    {"typ": "neue_seite", "querformat": True},
    {"typ": "absatz", "text": "quer"}]})
FOLIEN = lesen({"art": "folien", "titel": "Thema", "folien": [
    {"titel": "Eins", "inhalt": [{"typ": "text", "text": "Hallo"}]}]})


def als_vorlage(quelle: Path, ziel: Path, alt: str, neu: str) -> Path:
    """Gegenstück zu umwandeln(): macht aus .docx/.pptx eine .dotx/.potx, wie Word sie speichert."""
    with zipfile.ZipFile(quelle) as rein, zipfile.ZipFile(ziel, "w") as raus:
        for e in rein.infolist():
            daten = rein.read(e.filename)
            raus.writestr(e, daten.replace(alt.encode(), neu.encode()) if e.filename == "[Content_Types].xml"
                          else daten)
    return ziel


class TestVorlagen(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def word_vorlage(self, ohne_stile=False) -> Path:
        d = docx.Document()
        d.sections[0].header.paragraphs[0].text = "Schule Musterstadt"
        d.sections[0].left_margin = d.sections[0].right_margin = Cm(3.0)  # Textbreite 15 cm statt 17
        d.add_paragraph("Beispieltext, der verschwinden muss")
        if ohne_stile:
            for name in ("List Bullet", "List Number", "Heading 1", "Heading 2", "Heading 3"):
                el = d.styles[name].element
                el.getparent().remove(el)
        d.save(self.tmp / "vorlage.docx")
        return self.tmp / "vorlage.docx"

    def powerpoint_vorlage(self, format_=vorlagen.FOLIE_16_9) -> Path:
        prs = pptx.Presentation()
        prs.slide_width, prs.slide_height = (Emu(x) for x in format_)
        prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = "Musterfolie"
        prs.save(self.tmp / "vorlage.pptx")
        return self.tmp / "vorlage.pptx"

    def test_umwandeln_dotx_und_potx(self):
        dotx = als_vorlage(self.word_vorlage(), self.tmp / "v.dotx", "document.main+xml", "template.main+xml")
        docx_datei = vorlage_umwandeln(dotx, self.tmp / "v.docx")
        self.assertEqual(docx.Document(docx_datei).sections[0].header.paragraphs[0].text, "Schule Musterstadt")
        potx = als_vorlage(self.powerpoint_vorlage(), self.tmp / "v.potx", "presentation.main+xml",
                           "template.main+xml")
        pptx_datei = vorlage_umwandeln(potx, self.tmp / "v.pptx")
        self.assertEqual(len(pptx.Presentation(pptx_datei).slides), 1)

    def test_umwandeln_fehler(self):
        with self.assertRaisesRegex(ValueError, "nur als .dotx, .potx, .docx oder .pptx"):
            vorlage_umwandeln(self.tmp / "x.doc", self.tmp / "x.docx")
        with self.assertRaisesRegex(ValueError, "Ziel muss auf .docx enden"):
            vorlage_umwandeln(self.word_vorlage(), self.tmp / "x.pptx")
        falsch = self.tmp / "falsch.dotx"
        with zipfile.ZipFile(falsch, "w") as z:
            z.writestr("[Content_Types].xml", "<Types/>")
        with self.assertRaisesRegex(ValueError, "keine gültige .dotx-Datei"):
            vorlage_umwandeln(falsch, self.tmp / "f.docx")

    def test_word_vorlage_als_grundlage(self):
        ziel = arbeitsblatt.bauen(BLATT, self.tmp / "blatt.docx", vorlage=self.word_vorlage())
        d = docx.Document(ziel)
        self.assertEqual(d.sections[0].header.paragraphs[0].text, "Schule Musterstadt")
        self.assertEqual(round(d.sections[0].left_margin.cm, 2), 3.0)
        texte = [p.text for p in d.paragraphs]
        self.assertNotIn("Beispieltext, der verschwinden muss", texte)
        self.assertIn("Teil 1", texte)
        s = d.sections[0]
        textbreite = (s.page_width - s.left_margin - s.right_margin) / 360000  # cm, schmaler als die 16 cm des Kopfs
        self.assertAlmostEqual(sum(c.width.cm for c in d.tables[0].columns), textbreite, places=1)
        quer = d.sections[1]
        self.assertGreater(quer.page_width, quer.page_height)
        self.assertEqual(round(quer.left_margin.cm, 2), 3.0)

    def test_word_vorlage_ohne_stile(self):
        ziel = arbeitsblatt.bauen(BLATT, self.tmp / "blatt.docx", vorlage=self.word_vorlage(ohne_stile=True))
        texte = [p.text for p in docx.Document(ziel).paragraphs]
        self.assertTrue({"1.\ta", "2.\tb", "•\tc", "Teil 1"} <= set(texte))

    def test_powerpoint_vorlage_als_grundlage(self):
        vorlage = self.powerpoint_vorlage()
        master = len(pptx.Presentation(vorlage).slide_master.shapes)
        ziel = folien.bauen(FOLIEN, self.tmp / "f.pptx", vorlage=vorlage)
        prs = pptx.Presentation(ziel)
        self.assertEqual(len(prs.slides), 2)  # Titelfolie und eine Folie, ohne die Musterfolie
        self.assertNotIn("Musterfolie", [s.text_frame.text for f in prs.slides for s in f.shapes if s.has_text_frame])
        self.assertEqual(len(prs.slide_master.shapes), master)
        self.assertEqual(prs.slides[1].shapes.title.text_frame.text, "Eins")

    def test_powerpoint_vorlage_nicht_16_9(self):
        with self.assertRaisesRegex(ValueError, "nötig ist 16:9"):
            folien.bauen(FOLIEN, self.tmp / "f.pptx", vorlage=self.powerpoint_vorlage((9144000, 6858000)))

    def test_layout_nach_platzhaltern(self):
        prs = pptx.Presentation()
        self.assertEqual(vorlagen.layout(prs, nur_titel=True).name, "Title Only")
        self.assertEqual(vorlagen.layout(prs, nur_titel=False).name, "Blank")

    def test_falsche_vorlage(self):
        with self.assertRaisesRegex(ValueError, "Vorlage muss .docx sein"):
            arbeitsblatt.bauen(BLATT, self.tmp / "b.docx", vorlage=self.tmp / "v.dotx")
        with self.assertRaisesRegex(ValueError, "Vorlage fehlt"):
            arbeitsblatt.bauen(BLATT, self.tmp / "b.docx", vorlage=self.tmp / "fehlt.docx")


if __name__ == "__main__":
    unittest.main()
