"""Folien als pptx: bauen, wieder öffnen, Inhalt, Lage und Fehler prüfen. Das Aussehen prüft die Vorschau."""

import tempfile
import unittest
from pathlib import Path

import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

from renderer import BeschreibungUngueltig, folien
from renderer.pruefung import lesen

PROFIL = {"farben": {"akzent": "C00000", "kasten": "F5F5F5", "satzanfang": "196B24"}}
LANG = ("Yasmins großer Bruder zieht heute um. Die beiden verstehen sich sehr gut, er hilft ihr oft bei den "
        "Hausaufgaben und ist immer für sie da. Jetzt braucht er den ganzen Tag Hilfe beim Tragen seiner Möbel. ")

FOLIEN = {
    "art": "folien",
    "titel": "Muss ich helfen?",
    "untertitel": "Ethik · Klasse 8",
    "icon": "handschlag",
    "folien": [
        {"titel": "Argument", "inhalt": [
            {"typ": "kasten", "titel": "Ein Argument:", "absaetze": [[
                {"text": "Ich finde.", "farbe": "blau", "unterstrichen": True}, {"text": " "},
                {"text": "Weil.", "farbe": "orange", "unterstrichen": True}]]},
            {"typ": "legende", "eintraege": [{"text": "Behauptung", "farbe": "blau"},
                                             {"text": "Begründung", "farbe": "orange"}]}]},
        {"titel": "Ziel für heute", "icons": ["ziel"], "inhalt": [{"typ": "zentriert", "absaetze": ["Wir lernen."]}]},
        {"titel": "Yasmins Entscheidung", "icons": ["paket", "schale"], "inhalt": [
            {"typ": "kasten", "absaetze": [LANG * 3, [{"text": "Sie muss sich entscheiden.", "fett": True}]]},
            {"typ": "satzanfang", "text": "Yasmin steht vor dem Problem, dass …"}]},
        {"titel": "Unsere Leitfrage", "inhalt": [{"typ": "leitfrage", "text": "Muss ich helfen?"}]},
        {"titel": "Merksätze", "inhalt": [{"typ": "spalten", "spalten": [{"titel": "Singer", "text": "Not groß."},
                                                                           {"titel": "Aristoteles", "text": "Nähe."}]}]},
        {"titel": "Arbeitsauftrag", "inhalt": [{"typ": "zentriert", "absaetze": ["Bearbeitet 1–4!"]},
                                               {"typ": "zeit", "minuten": 10}]},
        {"titel": "Wer überzeugt?", "inhalt": [{"typ": "auswahl", "optionen": ["Aristoteles", "Singer"]},
                                               {"typ": "satzanfang", "text": "„Mich überzeugt …“"}]},
        {"titel": "Feedback", "inhalt": [{"typ": "feedback", "gut": ["Gelernt habe ich …"],
                                          "schlecht": ["Nicht verstanden habe ich …"]}]},
        {"titel": "Zum Schluss", "inhalt": [{"typ": "text", "text": "Danke."}, {"typ": "punkte", "punkte": ["a", "b"]}]},
    ],
}


def bauen(beschreibung=FOLIEN, profil=PROFIL):
    ziel = Path(tempfile.mkdtemp()) / "folien.pptx"
    folien.bauen(lesen(beschreibung), ziel, profil)
    return pptx.Presentation(ziel)


def texte(folie) -> list[str]:
    return [s.text_frame.text for s in folie.shapes if s.has_text_frame and s.text_frame.text]


def runs(folie):
    return [r for s in folie.shapes if s.has_text_frame for p in s.text_frame.paragraphs for r in p.runs]


def cm(emu) -> float:
    return round(emu / 360000, 2)


class TestFolien(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.prs = bauen()
        cls.f = cls.prs.slides

    def test_format_und_anzahl(self):
        self.assertEqual((cm(self.prs.slide_width), cm(self.prs.slide_height)), (33.87, 19.05))
        self.assertEqual(len(self.f), 10)
        self.assertEqual(self.prs.core_properties.author, "")

    def test_titelfolie(self):
        t = self.f[0]
        self.assertEqual(texte(t), ["Muss ich helfen?", "Ethik · Klasse 8"])
        arten = [s.shape_type for s in t.shapes]
        self.assertEqual(arten.count(MSO_SHAPE_TYPE.PICTURE), 1)
        balken = [s for s in t.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE][0]
        self.assertEqual(str(balken.fill.fore_color.rgb), "C00000")
        self.assertEqual((cm(balken.width), cm(balken.height)), (10.16, 19.05))

    def test_inhaltsfolie_rahmen(self):
        f = self.f[1]
        self.assertEqual(f.shapes.title.text_frame.text, "Argument")  # im Titel-Platzhalter: Gliederung
        titel_run = f.shapes.title.text_frame.paragraphs[0].runs[0]
        self.assertEqual((str(titel_run.font.color.rgb), titel_run.font.bold), ("C00000", True))
        self.assertIn("Ethik · Klasse 8 · Muss ich helfen?", texte(f))
        self.assertIn("2", texte(f))  # Foliennummer
        strich = [s for s in f.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE and cm(s.height) == 0.07]
        self.assertEqual(len(strich), 1)

    def test_icons_rechts_oben_in_reihenfolge(self):
        bilder = [s for s in self.f[3].shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE and cm(s.top) < 4.7]
        self.assertEqual(len(bilder), 2)
        self.assertLess(bilder[0].left, bilder[1].left)
        self.assertAlmostEqual(cm(bilder[1].left + bilder[1].width), folien.X1, places=1)

    def test_kasten_und_legende(self):
        r = {x.text: x for x in runs(self.f[1])}
        self.assertTrue(r["Ein Argument:"].font.bold)
        self.assertEqual((str(r["Ich finde."].font.color.rgb), r["Ich finde."].font.underline), ("156082", True))
        self.assertEqual(str(r["Weil."].font.color.rgb), "A8480C")
        self.assertTrue(r["Behauptung"].font.bold and r["Behauptung"].font.underline)

    def test_satzanfang_mit_anfuehrungszeichen(self):
        self.assertIn("„Yasmin steht vor dem Problem, dass …“", texte(self.f[3]))
        self.assertIn("„Mich überzeugt …“", texte(self.f[7]))  # nicht doppelt
        r = [x for x in runs(self.f[3]) if x.text.startswith("„")][0]
        self.assertEqual(str(r.font.color.rgb), "196B24")

    def test_weitere_elemente(self):
        self.assertIn("Muss ich helfen?", texte(self.f[4]))
        self.assertTrue({"Singer", "Not groß.", "Aristoteles", "Nähe."} <= set(texte(self.f[5])))
        self.assertIn("Zeit: 10 Minuten", texte(self.f[6]))
        self.assertTrue({"Aristoteles", "Singer"} <= set(texte(self.f[7])))
        self.assertTrue({"„Gelernt habe ich …“", "„Nicht verstanden habe ich …“"} <= set(texte(self.f[8])))
        self.assertEqual(sum(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in self.f[8].shapes), 2)  # Smileys

    def test_aufzaehlung(self):
        punkte = [s for s in self.f[9].shapes if s.has_text_frame and s.text_frame.text == "a\nb"][0]
        for p in punkte.text_frame.paragraphs:
            self.assertEqual(p._p.pPr.find(qn("a:buChar")).get("char"), "•")

    def test_nichts_ragt_in_die_fusszeile(self):
        for nr, f in enumerate(self.f):
            if nr == 0:
                continue
            for s in f.shapes:
                if s.has_text_frame and cm(s.top) >= 17.9:  # Fußzeile selbst
                    continue
                self.assertLessEqual(cm(s.top + s.height), folien.Y1 + 0.01, f"Folie {nr + 1}: {s.name}")

    def test_keine_schatten_aus_dem_design(self):
        for f in self.f:
            for s in f.shapes:
                if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE:
                    self.assertIsNone(s._element.find(qn("p:style")), s.name)

    def test_wenig_text_wird_groesser(self):
        kurz = [x for x in runs(self.f[1]) if x.text == "Ich finde."][0].font.size.pt
        lang = [x for x in runs(self.f[3]) if x.text.startswith("Yasmins großer")][0].font.size.pt
        self.assertGreater(kurz, lang)

    def test_ohne_profil_und_ohne_icon(self):
        d = dict(FOLIEN, folien=FOLIEN["folien"][:1])
        del d["icon"], d["untertitel"]
        prs = bauen(d, profil=None)
        self.assertEqual(sum(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in prs.slides[0].shapes), 0)
        self.assertIn("Muss ich helfen?", texte(prs.slides[1]))  # Fußzeile nur mit Titel


class TestUeberlauf(unittest.TestCase):
    def test_zu_viel_inhalt(self):
        d = dict(FOLIEN, folien=[{"titel": "Viel", "inhalt": [{"typ": "kasten", "absaetze": [LANG * 3] * 4}]},
                                 FOLIEN["folien"][1]])
        with self.assertRaises(BeschreibungUngueltig) as ctx:
            bauen(d)
        self.assertEqual(ctx.exception.fehler, ["folien[0].inhalt: zu viel für eine Folie, auch verkleinert. "
                                                "Kürzen oder auf zwei Folien verteilen."])

    def test_titel_zu_lang(self):
        d = dict(FOLIEN, folien=[{"titel": "Sehr langer Titel " * 10, "inhalt": [{"typ": "text", "text": "x"}]}])
        with self.assertRaises(BeschreibungUngueltig) as ctx:
            bauen(d)
        self.assertIn("folien[0].titel: zu lang für die Titelzeile", ctx.exception.fehler[0])


class TestZeilen(unittest.TestCase):
    def test_umbruch(self):
        self.assertEqual(folien.zeilen("kurz", 10, 20), 1)
        self.assertEqual(folien.zeilen("a\nb", 10, 20), 2)
        self.assertGreater(folien.zeilen("Wort " * 40, 10, 20), 5)
        self.assertGreater(folien.zeilen("x" * 200, 5, 20), 1)  # überlanges Wort bricht
        self.assertGreaterEqual(folien.zeilen("Wort " * 20, 10, 20, fett=True), folien.zeilen("Wort " * 20, 10, 20))


if __name__ == "__main__":
    unittest.main()
