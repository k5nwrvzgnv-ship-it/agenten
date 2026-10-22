"""Zusammenfassung: Prüfung (Fehler mit Pfad für Claude) und Bauen als docx (Inhalt, Aufbau, Vorlage)."""

import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import docx
from docx.oxml.ns import qn
from PIL import Image

from renderer import bauen, design, pruefen, zusammenfassung
from renderer.pruefung import lesen

BEISPIEL = Path(__file__).resolve().parent.parent / "beispiele" / "zusammenfassung-beispiel.json"

ZF = {
    "art": "zusammenfassung",
    "titel": "Grundlagen",
    "kurs": "Holzschutz",
    "datum": "20.10.2026",
    "bloecke": [
        {"typ": "ueberschrift", "text": "Erstens"},
        {"typ": "absatz", "text": "Satz.", "quellen": [{"art": "transkript", "stelle": "2-3"},
                                                      {"art": "folie", "stelle": "4"}]},
        {"typ": "ueberschrift", "text": "Unterpunkt", "ebene": 2},
        {"typ": "liste", "nummeriert": True, "punkte": [
            "a", {"text": "b", "kennzeichnung": "unsicher", "quellen": [{"art": "norm", "stelle": "DIN 68800-1"}]}]},
        {"typ": "ueberschrift", "text": "Zweitens"},
        {"typ": "ueberschrift", "text": "Wieder eins", "ebene": 2},
        {"typ": "hinweis", "art": "pruefungsrelevant", "text": "Kommt dran.", "kennzeichnung": "muendlich"},
        {"typ": "tabelle", "titel": "Klassen", "kopf": ["GK", "Lage"], "zeilen": [["1", "innen"]],
         "quellen": [{"art": "folie", "stelle": "9"}]},
        {"typ": "rechenbeispiel", "aufgabe": "Rechne.", "gegeben": ["a = 1"],
         "schritte": [{"text": "Einsetzen", "rechnung": "1 + 1"}, {"text": "Fertig"}], "ergebnis": "2"},
    ],
    "selbsttest": [{"frage": "Was?", "loesung": "Das.", "quellen": [{"art": "transkript", "stelle": "5"}]}],
    "offene_fragen": ["Warum?"],
}
PROFIL = {"erstellt_von": "Max Muster", "farben": {"ueberschrift": "1E5A3C", "pruefungsrelevant": "DDEEDD"}}


def mit(**aenderungen) -> dict:
    """Kopie von ZF, z. B. mit(b1={...}) ersetzt bloecke[1], mit(datum=...) ein Feld oben."""
    z = copy.deepcopy(ZF)
    for name, wert in aenderungen.items():
        if name[0] == "b" and name[1:].isdigit():
            z["bloecke"][int(name[1:])] = wert
        else:
            z[name] = wert
    return z


class TestPruefung(unittest.TestCase):
    def test_gueltig_mit_standardwerten(self):
        b = lesen({"art": "zusammenfassung", "titel": "T", "kurs": "K", "datum": "01.02.2026",
                   "bloecke": [{"typ": "liste", "punkte": ["a"]}]})
        self.assertEqual(b, {"art": "zusammenfassung", "titel": "T", "untertitel": "", "kurs": "K",
                             "datum": "01.02.2026", "selbsttest": [], "offene_fragen": [], "bloecke": [
                                 {"typ": "liste", "nummeriert": False,
                                  "punkte": [{"text": "a", "quellen": [], "kennzeichnung": None}]}]})

    def test_beispiele_gueltig(self):
        self.assertEqual(pruefen(ZF), [])
        self.assertEqual(pruefen(BEISPIEL.read_text(encoding="utf-8")), [])

    def test_kopf_fehler(self):
        z = mit(datum="31.02.2026")
        del z["kurs"]
        self.assertEqual(pruefen(z), [
            "kurs: fehlt", 'datum: muss ein Datum wie "20.10.2026" sein (TT.MM.JJJJ), ist "31.02.2026"'])
        self.assertIn('datum: muss ein Datum wie "20.10.2026" sein (TT.MM.JJJJ), ist "2026-10-20"',
                      pruefen(mit(datum="2026-10-20")))

    def test_bloecke_des_arbeitsblatts_nicht_erlaubt(self):
        f = pruefen(mit(b0={"typ": "merksatz", "text": "x"}))
        self.assertEqual(f, ['bloecke[0].typ: muss einer von "ueberschrift", "absatz", "liste", "tabelle", '
                             '"hinweis", "rechenbeispiel" sein, ist "merksatz"'])

    def test_ueberschriften(self):
        self.assertEqual(pruefen(mit(b0={"typ": "ueberschrift", "text": "1.2 Erstens"})),
                         ["bloecke[0].text: ohne Nummer schreiben, die Nummer setzt der Renderer"])
        self.assertEqual(pruefen(mit(b0={"typ": "ueberschrift", "text": "Erstens", "ebene": 3})),
                         ["bloecke[0].ebene: muss eine ganze Zahl von 1 bis 2 sein, ist 3"])
        self.assertEqual(pruefen(mit(b0={"typ": "ueberschrift", "text": "Erstens", "ebene": 2})),
                         [f"bloecke[{i}].ebene: Ebene 2 vor der ersten Überschrift der Ebene 1, "
                          "erst eine Überschrift mit ebene 1 setzen" for i in (0, 2)])
        self.assertEqual(pruefen(mit(b0={"typ": "ueberschrift", "text": "1. Erstens"})),
                         ["bloecke[0].text: ohne Nummer schreiben, die Nummer setzt der Renderer"])
        for text in ("2020 war ein Jahr", "3 Säulen des Holzschutzes", "1950er Jahre"):
            self.assertEqual(pruefen(mit(b0={"typ": "ueberschrift", "text": text})), [], text)

    def test_quellen_und_kennzeichnung(self):
        f = pruefen(mit(b1={"typ": "absatz", "text": "Satz [Transkript S. 2].", "kennzeichnung": "unsicherr",
                            "quellen": [{"art": "folie", "stelle": "S. 4"}, {"art": "buch", "stelle": "x"},
                                        {"art": "norm"}, "Folie 3"]}))
        self.assertEqual(f, [
            'bloecke[1].text: Quellenangabe im Text, bitte ins Feld quellen verschieben '
            '(z. B. [{"art": "transkript", "stelle": "4"}])',
            'bloecke[1].kennzeichnung: muss "unsicher" oder "muendlich" sein, ist "unsicherr" (gemeint: "unsicher"?)',
            'bloecke[1].quellen[0].stelle: nur Foliennummern wie "4", "4-5" oder "4, 7", ist "S. 4"',
            'bloecke[1].quellen[1].art: muss einer von "transkript", "folie", "norm" sein, ist "buch"',
            "bloecke[1].quellen[2].stelle: fehlt",
            'bloecke[1].quellen[3]: muss ein Objekt {"art": ..., "stelle": ...} sein, ist "Folie 3"',
        ])
        self.assertEqual(pruefen(mit(b1={"typ": "absatz", "text": "x", "quellen": []})),
                         ["bloecke[1].quellen: muss 1 bis 10 Einträge haben, hat 0"])
        self.assertEqual(pruefen(mit(b1={"typ": "absatz", "text": "x", "quellen": [
            {"art": "transkript", "stelle": "4, 7"}, {"art": "transkript", "stelle": "4–5"}]})), [])

    def test_liste_und_offene_fragen(self):
        f = pruefen(mit(b3={"typ": "liste", "punkte": [3, {"txt": "a"}]}, offene_fragen=[{"text": "x", "farbe": 1}]))
        self.assertEqual(f, [
            'bloecke[3].punkte[0]: muss ein Text oder ein Objekt {"text": ...} sein, ist 3',
            'bloecke[3].punkte[1].txt: unbekanntes Feld (gemeint: "text"?), erlaubt sind: text, quellen, kennzeichnung',
            "bloecke[3].punkte[1].text: fehlt",
            "offene_fragen[0].farbe: unbekanntes Feld, erlaubt sind: text, quellen, kennzeichnung",
        ])

    def test_tabelle_hinweis_rechenbeispiel(self):
        f = pruefen(mit(b6={"typ": "hinweis", "art": "wichtig", "text": "x"},
                        b7={"typ": "tabelle", "zeilen": [["a", "b"], ["c"]], "zeilenhoehe_cm": 2},
                        b8={"typ": "rechenbeispiel", "aufgabe": "x", "schritte": [{"rechnung": "1"}, "2"]}))
        self.assertEqual(f, [
            'bloecke[6].art: muss "pruefungsrelevant" oder "organisatorisch" sein, ist "wichtig"',
            "bloecke[7].zeilenhoehe_cm: unbekanntes Feld (gemeint: \"zeilen\"?), erlaubt sind: typ, zeilen, titel, "
            "kopf, erste_spalte_fett, quellen, kennzeichnung",
            "bloecke[7].zeilen[1]: hat 1 Zellen, erwartet 2 (wie zeilen[0])",
            "bloecke[8].ergebnis: fehlt",
            "bloecke[8].schritte[0].text: fehlt",
            'bloecke[8].schritte[1]: muss ein Objekt {...} sein, ist "2"',
        ])

    def test_selbsttest(self):
        f = pruefen(mit(selbsttest=[{"frage": "x"}, {"frage": "y", "loesung": "z", "kennzeichnung": "unsicher"}]))
        self.assertEqual(f, ["selbsttest[0].loesung: fehlt",
                             "selbsttest[1].kennzeichnung: unbekanntes Feld, erlaubt sind: frage, loesung, quellen"])
        self.assertEqual(pruefen(mit(selbsttest=[])), ["selbsttest: muss 1 bis 30 Einträge haben, hat 0"])


class TestBauen(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        heute = mock.patch.object(zusammenfassung, "heute", return_value="22.10.2026")
        heute.start()
        self.addCleanup(heute.stop)

    def bauen(self, beschreibung=ZF, profil=PROFIL, vorlage=None):
        return bauen(beschreibung, self.tmp / "z.docx", profil, vorlage=vorlage, vorschau=False)

    def test_kontrolltext(self):
        self.assertEqual(self.bauen().kontrolltext.split("\n"), [
            "Holzschutz",
            "Grundlagen",
            "Vorlesung vom 20.10.2026 · Stand 22.10.2026 · erstellt von Max Muster",
            "[unsicher] Aussage im Transkript nicht eindeutig, bitte prüfen   "
            "[nur mündlich] nur in der Vorlesung gesagt, nicht auf den Folien",
            "# 1  Erstens",
            "Satz. [Transkript S. 2-3; Folie 4]",
            "## 1.1  Unterpunkt",
            "1. a",
            "2. b [unsicher] [DIN 68800-1]",
            "# 2  Zweitens",
            "## 2.1  Wieder eins",
            "[Kasten] Prüfungsrelevant",
            "[Kasten] Kommt dran. [nur mündlich]",
            "Klassen",
            "| GK | Lage |",
            "| 1 | innen |",
            "[Folie 9]",
            "[Kasten] Rechenbeispiel",
            "[Kasten] Rechne.",
            "[Kasten] Gegeben:",
            "[Kasten] •\ta = 1",
            "[Kasten] Schritt 1: Einsetzen",
            "[Kasten] \t1 + 1",
            "[Kasten] Schritt 2: Fertig",
            "[Kasten] Ergebnis: 2",
            "# Selbsttest",
            "1.\tWas?",
            "# Offene Fragen",
            "• Warum?",
            "--- neue Seite",
            "# Lösungen zum Selbsttest",
            "1.\tDas. [Transkript S. 5]",
        ])

    def test_ohne_kennzeichen_und_autor_kein_hinweis(self):
        z = {"art": "zusammenfassung", "titel": "T", "kurs": "K", "datum": "01.02.2026",
             "bloecke": [{"typ": "absatz", "text": "x"}]}
        self.assertEqual(self.bauen(z, profil=None).kontrolltext.split("\n"),
                         ["K", "T", "Vorlesung vom 01.02.2026 · Stand 22.10.2026", "x"])

    def test_farben_aus_dem_profil(self):
        d = docx.Document(self.bauen().pfad)
        self.assertEqual(str(d.styles["Heading 1"].font.color.rgb), "1E5A3C")
        flaechen = {el.get(qn("w:fill")) for el in d.element.body.iter(qn("w:shd"))}
        self.assertIn("DDEEDD", flaechen)  # Hinweis prüfungsrelevant
        self.assertIn(design.STANDARD["farben"]["tabellenkopf"], flaechen)
        self.assertIn(design.STANDARD["farben"]["rechenbeispiel"], flaechen)
        self.assertEqual(d.core_properties.author, "Max Muster")

    def test_fusszeile_mit_seitenzahlen(self):
        d = docx.Document(self.bauen().pfad)
        fuss = d.sections[0].footer
        self.assertEqual(fuss.paragraphs[0].text, "Holzschutz · Grundlagen\tSeite 1 von 1")
        befehle = [el.text.strip() for el in fuss._element.iter(qn("w:instrText"))]
        self.assertEqual(befehle, ["PAGE", "NUMPAGES"])

    def test_vorlage_mit_logo(self):
        logo = self.tmp / "logo.png"
        Image.new("RGB", (40, 20), "green").save(logo)
        v = docx.Document()
        v.sections[0].header.paragraphs[0].add_run().add_picture(str(logo))
        v.add_paragraph("weg damit")
        v.save(self.tmp / "vorlage.docx")
        e = self.bauen(vorlage=self.tmp / "vorlage.docx")
        d = docx.Document(e.pfad)
        self.assertEqual(len(d.sections[0].header._element.findall(".//" + qn("w:drawing"))), 1)
        self.assertNotIn("weg damit", e.kontrolltext)
        self.assertIn("Seite", d.sections[0].footer.paragraphs[0].text)  # Vorlage ohne Fußzeile: eigene

    def test_fusszeile_der_vorlage_bleibt(self):
        v = docx.Document()
        v.sections[0].footer.paragraphs[0].text = "Hochschule"
        v.save(self.tmp / "vorlage.docx")
        d = docx.Document(self.bauen(vorlage=self.tmp / "vorlage.docx").pfad)
        self.assertEqual([p.text for p in d.sections[0].footer.paragraphs], ["Hochschule"])

    def test_vorlage_ohne_listenstile(self):
        v = docx.Document()
        for name in ("List Bullet", "List Number"):
            el = v.styles[name].element
            el.getparent().remove(el)
        v.save(self.tmp / "vorlage.docx")
        text = self.bauen(vorlage=self.tmp / "vorlage.docx").kontrolltext
        self.assertIn("1.\ta", text)
        self.assertIn("•\tWarum?", text)

    def test_beispiel_baut(self):
        e = self.bauen(json.loads(BEISPIEL.read_text(encoding="utf-8")))
        self.assertIn("# 3  Holzfeuchte", e.kontrolltext)


class TestProfil(unittest.TestCase):
    def test_neue_felder(self):
        p = design.profil_lesen({"erstellt_von": "Jan", "farben": {"kennzeichnung": "#123abc"}})
        self.assertEqual((p["erstellt_von"], p["farben"]["kennzeichnung"]), ("Jan", "123ABC"))
        with self.assertRaisesRegex(design.ProfilUngueltig, "erstellt_von: muss ein Text"):
            design.profil_lesen({"erstellt_von": 5})


if __name__ == "__main__":
    unittest.main()
