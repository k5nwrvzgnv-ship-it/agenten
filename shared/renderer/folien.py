"""Folien als pptx bauen, aus einer geprüften Beschreibung (pruefung.lesen) und einem Design-Profil.

Aufbau nach Lias Vorbild (16:9): Titelfolie mit farbigem Balken und Icon, Inhaltsfolien mit Titel in der
Akzentfarbe, Strich darunter, optional Icons rechts oben, Fußzeile mit Thema und Nummer. Der Inhalt besteht
aus Bausteinen (Elementen), die untereinander stehen.

PowerPoint verkleinert Text, der nicht passt, erst beim Bearbeiten. Darum schätzt der Renderer die Höhe jedes
Elements aus Textlänge und Schriftgröße und wählt je Folie den größten Maßstab, bei dem alles passt. Passt es
auch verkleinert nicht, geht ein Fehler an Claude zurück. Die Schätzung ist bewusst vorsichtig; die Vorschau
zeigt, ob sie stimmt.
"""

import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Cm, Emu, Pt

from renderer import design, icons, vorlagen
from renderer.pruefung import BeschreibungUngueltig

FOLIE_CM = (33.867, 19.05)
X0, X1 = 2.37, 31.5  # Inhalt links und rechts (wie Lias Titelstrich)
BREITE = X1 - X0
Y0, Y1 = 5.5, 17.6  # Inhaltsbereich unter dem Titelstrich, über der Fußzeile
ABSTAND = 0.7  # zwischen Elementen
TITEL = {"y": 1.2, "hoehe": 3.4, "groessen": ((44, 1), (38, 2), (32, 2))}  # (Größe, höchstens Zeilen)
STRICH = {"y": 4.72, "hoehe": 0.07}
ICON_CM = 2.4
SCHWARZ, GRAU_FUSS, GRAU_UNTERTITEL, WEISS = "000000", "7F7F7F", "595959", "FFFFFF"
TEXTFARBEN = {"blau": "156082", "orange": "A8480C", "gruen": "196B24", "rot": "C00000", "grau": "595959"}

# Grundgröße (pt) je Element und wie stark es auf leeren Folien wachsen darf
GROESSE = {"text": 24, "punkte": 24, "kasten": 22, "leitfrage": 40, "zentriert": 30, "zeit": 28, "spalten": 22,
           "auswahl": 30, "satzanfang": 28, "feedback": 18, "legende": 24}
WACHSEN = {"text": 1.25, "punkte": 1.2, "kasten": 1.35}
MASSSTAEBE = [round(1.35 - 0.05 * i, 2) for i in range(16)]  # 1.35 bis 0.6
ZEILE = 1.2  # Zeilenhöhe in Vielfachen der Schriftgröße
MITTIG = {"zentriert", "zeit"}  # Folien nur aus diesen Elementen stehen senkrecht mittig


def bauen(beschreibung: dict, ziel: Path, profil: dict | None = None, vorlage: Path | None = None) -> Path:
    """Baut die Folien nach `ziel` (.pptx). `beschreibung` aus pruefung.lesen(), `profil` siehe design.py,
    `vorlage` eine .pptx in 16:9 (Master und Hintergrund). Wirft BeschreibungUngueltig, wenn der Inhalt einer
    Folie auch verkleinert nicht passt."""
    with tempfile.TemporaryDirectory(prefix="renderer-icons-") as tmp:
        bau = _Bau(beschreibung, design.profil_lesen(profil), Path(tmp), vorlage)
        bau.titelfolie()
        fehler = []
        for nr, folie in enumerate(beschreibung["folien"]):
            fehler += bau.inhaltsfolie(nr, folie)
        if fehler:
            raise BeschreibungUngueltig(fehler)
        bau.prs.save(ziel)
    return ziel


# Textmaß: grobe Zeichenbreiten in Vielfachen der Schriftgröße (vorsichtig, eher breiter als Aptos)

def _breite_zeichen(z: str) -> float:
    if z in "iljtfrI.,;:!'|() ":
        return 0.3
    if z in "mwMW":
        return 0.85
    if z.isupper() or z.isdigit():
        return 0.64
    return 0.55


def zeilen(text: str, breite_cm: float, groesse_pt: float, fett: bool = False) -> int:
    """Wie viele Zeilen `text` bei dieser Breite braucht (Umbruch an Leerzeichen, wie PowerPoint)."""
    breite_pt = breite_cm / 2.54 * 72
    faktor = groesse_pt * (1.06 if fett else 1.0)
    anzahl = 0
    for absatz in text.split("\n"):
        anzahl += 1
        rest = breite_pt
        for wort in absatz.split(" "):
            w = sum(_breite_zeichen(z) for z in wort) * faktor
            leer = _breite_zeichen(" ") * faktor
            if w + leer <= rest or rest == breite_pt:
                rest -= w + leer
                while rest < 0:  # Wort länger als eine Zeile
                    anzahl += 1
                    rest += breite_pt
            else:
                anzahl += 1
                rest = breite_pt - w - leer
    return anzahl


def _cm(pt: float) -> float:
    return pt / 72 * 2.54


def _teile(text: str) -> list:
    return [{"text": text, "fett": False, "farbe": None, "unterstrichen": False}]


def _klartext(teile: list) -> str:
    return "".join(t["text"] for t in teile)


def _zitat(text: str) -> str:
    """Satzanfänge stehen in deutschen Anführungszeichen."""
    return text if text.startswith(("„", '"')) else f"„{text}“"


def _texthoehe(absaetze: list, breite: float, groesse: float, fett=False, abstand_pt: float = 0) -> float:
    n = sum(zeilen(_klartext(a), breite, groesse, fett or any(t["fett"] for t in a)) for a in absaetze)
    return _cm(n * groesse * ZEILE + abstand_pt * (len(absaetze) - 1))


class _Bau:
    def __init__(self, beschreibung: dict, profil: dict, tmp: Path, vorlage: Path | None):
        self.b, self.profil, self.tmp = beschreibung, profil, tmp
        self.farben = profil["farben"]
        if vorlage:
            self.prs = vorlagen.powerpoint_grundlage(vorlage)
        else:
            self.prs = Presentation()
            self.prs.slide_width, self.prs.slide_height = (Emu(x) for x in vorlagen.FOLIE_16_9)
        self.layout_leer = vorlagen.layout(self.prs, nur_titel=False)
        self.layout_titel = vorlagen.layout(self.prs, nur_titel=True)
        self.prs.core_properties.title = beschreibung["titel"]
        self.prs.core_properties.author = ""  # sonst "python-pptx"
        self.prs.core_properties.last_modified_by = ""
        self.icon_nr = 0

    # Grundformen

    def _textfeld(self, folie, x, y, w, h, absaetze, groesse, fett=False, farbe=SCHWARZ, ausrichtung=PP_ALIGN.LEFT,
                  anker=MSO_ANCHOR.TOP, abstand_pt=0, schrift=None, form=None):
        """Textfeld (oder Text in `form`) mit Absätzen aus Textteilen, ohne Innenrand und ohne Autoanpassung."""
        tb = form or folie.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
        tf = tb.text_frame
        tf.word_wrap, tf.auto_size, tf.vertical_anchor = True, MSO_AUTO_SIZE.NONE, anker
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, teile in enumerate(absaetze):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment, p.line_spacing = ausrichtung, 1.0  # einfacher Zeilenabstand, etwa ZEILE x Größe
            if i < len(absaetze) - 1:
                p.space_after = Pt(abstand_pt)
            for t in teile:
                run = p.add_run()
                run.text = t["text"]
                f = run.font
                f.name, f.size = schrift or self.profil["folien_schrift"], Pt(groesse)
                f.bold, f.underline = fett or t["fett"], t["unterstrichen"]
                f.color.rgb = RGBColor.from_string(TEXTFARBEN.get(t["farbe"], farbe))
        return tb

    def _flaeche(self, folie, x, y, w, h, fuellung, rund=False, rand=None):
        form = folie.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rund else MSO_SHAPE.RECTANGLE,
                                      Cm(x), Cm(y), Cm(w), Cm(h))
        if rund:
            form.adjustments[0] = min(0.5, 0.3 / h)  # Eckenradius etwa 0,3 cm
        form.fill.solid()
        form.fill.fore_color.rgb = RGBColor.from_string(fuellung)
        if rand:
            form.line.color.rgb = RGBColor.from_string(rand)
            form.line.width = Pt(1.5)
        else:
            form.line.fill.background()
        # Ohne Verweis auf die Stile des Designs: Dessen Effektstil zeichnet sonst in LibreOffice einen Schatten.
        stil = form._element.find(qn("p:style"))
        if stil is not None:
            form._element.remove(stil)
        return form

    def _icon(self, folie, name, x, y, groesse, farbe, gefuellt=False):
        self.icon_nr += 1
        pfad = icons.zeichnen(name, self.tmp / f"icon-{self.icon_nr}.png", farbe=farbe, gefuellt=gefuellt)
        bild = folie.shapes.add_picture(str(pfad), Cm(x), Cm(y), Cm(groesse), Cm(groesse))
        bild.name = f"Icon {name}"  # für den Kontrolltext und für Lia beim Bearbeiten

    # Folien

    def titelfolie(self):
        f = self.prs.slides.add_slide(self.layout_leer)
        self._flaeche(f, 0, 0, 10.16, FOLIE_CM[1], self.farben["akzent"])
        if self.b["icon"]:
            self._icon(f, self.b["icon"], (10.16 - 5.2) / 2, (FOLIE_CM[1] - 5.2) / 2, 5.2, WEISS)
        breite, titel = 20.0, self.b["titel"]
        groesse = next((g for g in (48, 40, 34, 28) if zeilen(titel, breite, g, True) * g * ZEILE <= 8.5 / 2.54 * 72),
                       28)
        self._textfeld(f, 12.1, 3.5, breite, 9.3, [_teile(titel)], groesse, fett=True, farbe=self.farben["akzent"],
                       anker=MSO_ANCHOR.BOTTOM, schrift=self.profil["folien_titelschrift"])
        if self.b["untertitel"]:
            self._textfeld(f, 12.1, 13.3, breite, 1.6, [_teile(self.b["untertitel"])], 24, farbe=GRAU_UNTERTITEL)

    def inhaltsfolie(self, nr: int, folie: dict) -> list[str]:
        """Baut eine Inhaltsfolie. Gibt Fehler zurück (Inhalt oder Titel passt nicht), sonst []."""
        f = self.prs.slides.add_slide(self.layout_titel)  # Titel im Platzhalter: steht in der Gliederung
        fehler = self._titel(f, nr, folie)
        self._flaeche(f, X0, STRICH["y"], BREITE, STRICH["hoehe"], self.farben["akzent"])
        anzahl = len(folie["icons"])
        for i, name in enumerate(folie["icons"]):  # von links nach rechts, rechtsbündig
            x = X1 - (anzahl - i) * ICON_CM - (anzahl - 1 - i) * 0.4
            self._icon(f, name, x, TITEL["y"] + (TITEL["hoehe"] - ICON_CM) / 2, ICON_CM, self.farben["akzent"],
                       gefuellt=True)
        self._fusszeile(f, nr + 2)
        elemente = folie["inhalt"]
        for skala in MASSSTAEBE:
            hoehen = [getattr(self, f"_hoehe_{e['typ']}")(e, self._groesse(e, skala)) for e in elemente]
            gesamt = sum(hoehen) + ABSTAND * skala * (len(elemente) - 1)
            if gesamt <= Y1 - Y0:
                break
        else:
            return fehler + [f"folien[{nr}].inhalt: zu viel für eine Folie, auch verkleinert. Kürzen oder auf zwei "
                             "Folien verteilen."]
        y = Y0 + ((Y1 - Y0 - gesamt) / 2 if all(e["typ"] in MITTIG for e in elemente) else 0)
        for e, h in zip(elemente, hoehen, strict=True):
            getattr(self, f"_zeichne_{e['typ']}")(f, e, y, h, self._groesse(e, skala))
            y += h + ABSTAND * skala
        return fehler

    def _groesse(self, e: dict, skala: float) -> int:
        return max(12, round(GROESSE[e["typ"]] * min(skala, WACHSEN.get(e["typ"], 1.0))))

    def _titel(self, f, nr: int, folie: dict) -> list[str]:
        breite = BREITE - len(folie["icons"]) * (ICON_CM + 0.4)
        groesse = next((g for g, n in TITEL["groessen"] if zeilen(folie["titel"], breite - 0.2, g, True) <= n), None)
        fehler = []
        if groesse is None:
            groesse = TITEL["groessen"][-1][0]
            fehler.append(f"folien[{nr}].titel: zu lang für die Titelzeile (höchstens zwei Zeilen, etwa 90 Zeichen)")
        form = f.shapes.title or f.shapes.add_textbox(0, 0, 0, 0)  # Vorlage ohne Titel-Layout
        form.left, form.top, form.width, form.height = Cm(X0), Cm(TITEL["y"]), Cm(breite), Cm(TITEL["hoehe"])
        self._textfeld(f, 0, 0, 0, 0, [_teile(folie["titel"])], groesse, fett=True, farbe=self.farben["akzent"],
                       anker=MSO_ANCHOR.MIDDLE, schrift=self.profil["folien_titelschrift"], form=form)
        form.text_frame.margin_left = Cm(0.21)
        return fehler

    def _fusszeile(self, f, nummer: int):
        u, t = self.b["untertitel"], self.b["titel"]
        text = f"{u} · {t}" if u else t
        self._textfeld(f, 2.33, 17.95, 26.0, 0.6, [_teile(text)], 10.8, farbe=GRAU_FUSS)
        self._textfeld(f, 29.0, 17.95, X1 - 29.0, 0.6, [_teile(str(nummer))], 10.8, farbe=GRAU_FUSS,
                       ausrichtung=PP_ALIGN.RIGHT)

    # Elemente: _hoehe_<typ>(e, groesse) -> cm, _zeichne_<typ>(folie, e, y, hoehe, groesse)

    def _hoehe_text(self, e, g):
        return _texthoehe([e["text"]], BREITE - 0.21, g)

    def _zeichne_text(self, f, e, y, h, g):
        self._textfeld(f, X0 + 0.21, y, BREITE - 0.21, h, [e["text"]], g)

    def _hoehe_punkte(self, e, g):
        return _texthoehe(e["punkte"], BREITE - 1.2, g, abstand_pt=8)

    def _zeichne_punkte(self, f, e, y, h, g):
        tb = self._textfeld(f, X0 + 0.21, y, BREITE - 0.21, h, e["punkte"], g, abstand_pt=8)
        for p in tb.text_frame.paragraphs:
            _xml_aufzaehlung(p, Cm(1.0))

    _KASTEN_INNEN = (3.22, 30.8)

    def _kasten_absaetze(self, e):
        return ([[{**_teile(e["titel"])[0], "fett": True}]] if e["titel"] else []) + e["absaetze"]

    def _hoehe_kasten(self, e, g):
        breite = self._KASTEN_INNEN[1] - self._KASTEN_INNEN[0]
        return 1.0 + _texthoehe(self._kasten_absaetze(e), breite, g, abstand_pt=g * 0.5)

    def _zeichne_kasten(self, f, e, y, h, g):
        self._flaeche(f, 2.54, y, X1 - 0.09 - 2.54, h, self.farben["kasten"], rund=True)
        self._flaeche(f, X0, y, 0.17, h, self.farben["akzent"])
        links, rechts = self._KASTEN_INNEN
        self._textfeld(f, links, y + 0.5, rechts - links, h - 1.0, self._kasten_absaetze(e), g, abstand_pt=g * 0.5)

    def _hoehe_leitfrage(self, e, g):
        return 2.0 + _texthoehe([_teile(e["text"])], BREITE - 2.0, g, fett=True)

    def _zeichne_leitfrage(self, f, e, y, h, g):
        form = self._flaeche(f, X0, y, BREITE, h, self.farben["akzent"], rund=True)
        self._textfeld(f, 0, 0, 0, 0, [_teile(e["text"])], g, fett=True, farbe=WEISS, ausrichtung=PP_ALIGN.CENTER,
                       anker=MSO_ANCHOR.MIDDLE, form=form)
        form.text_frame.margin_left = form.text_frame.margin_right = Cm(1.0)

    def _hoehe_zentriert(self, e, g):
        return _texthoehe(e["absaetze"], BREITE - 2.0, g, fett=True, abstand_pt=g * 0.5)

    def _zeichne_zentriert(self, f, e, y, h, g):
        self._textfeld(f, X0 + 1.0, y, BREITE - 2.0, h, e["absaetze"], g, fett=True, ausrichtung=PP_ALIGN.CENTER,
                       abstand_pt=g * 0.5)

    def _hoehe_zeit(self, e, g):
        return 1.9 + 0.25 + _cm(g * ZEILE)

    def _zeichne_zeit(self, f, e, y, h, g):
        self._icon(f, "wecker", (FOLIE_CM[0] - 1.9) / 2, y, 1.9, self.farben["akzent"])
        self._textfeld(f, X0, y + 2.15, BREITE, h - 2.15, [_teile(f"Zeit: {e['minuten']} Minuten")], g, fett=True,
                       farbe=self.farben["akzent"], ausrichtung=PP_ALIGN.CENTER)

    def _spalten_breite(self, e):
        n = len(e["spalten"])
        return (BREITE - 0.3 * (n - 1)) / n

    def _hoehe_spalten(self, e, g):
        innen = self._spalten_breite(e) - 0.8
        return 1.0 + max(_texthoehe([_teile(s["titel"])], innen, g + 2, fett=True) + 0.3
                         + _texthoehe([s["text"]], innen, g) for s in e["spalten"])

    def _zeichne_spalten(self, f, e, y, h, g):
        w = self._spalten_breite(e)
        for i, s in enumerate(e["spalten"]):
            x = X0 + i * (w + 0.3)
            self._flaeche(f, x, y, w, h, self.farben["kasten"])
            titel_h = _texthoehe([_teile(s["titel"])], w - 0.8, g + 2, fett=True)
            self._textfeld(f, x + 0.4, y + 0.5, w - 0.8, titel_h, [_teile(s["titel"])], g + 2, fett=True,
                           farbe=self.farben["akzent"])
            self._textfeld(f, x + 0.4, y + 0.5 + titel_h + 0.3, w - 0.8, h - 1.3 - titel_h, [s["text"]], g)

    def _auswahl_breite(self, e):
        n = len(e["optionen"])
        return (BREITE - 1.0 * (n - 1)) / n

    def _hoehe_auswahl(self, e, g):
        return 1.6 + max(_texthoehe([_teile(o)], self._auswahl_breite(e) - 1.0, g, fett=True) for o in e["optionen"])

    def _zeichne_auswahl(self, f, e, y, h, g):
        w = self._auswahl_breite(e)
        for i, option in enumerate(e["optionen"]):
            form = self._flaeche(f, X0 + i * (w + 1.0), y, w, h, self.farben["kasten"], rund=True,
                                 rand=self.farben["akzent"])
            self._textfeld(f, 0, 0, 0, 0, [_teile(option)], g, fett=True, farbe=self.farben["akzent"],
                           ausrichtung=PP_ALIGN.CENTER, anker=MSO_ANCHOR.MIDDLE, form=form)
            form.text_frame.margin_left = form.text_frame.margin_right = Cm(0.5)

    def _hoehe_satzanfang(self, e, g):
        return max(1.5, _texthoehe([_teile(_zitat(e["text"]))], BREITE - 2.75, g, fett=True))

    def _zeichne_satzanfang(self, f, e, y, h, g):
        self._icon(f, "zeigefinger", X0 + 0.3, y + (min(h, _cm(g * ZEILE) * 1.4) - 1.5) / 2, 1.5,
                   self.farben["satzanfang"])
        self._textfeld(f, X0 + 2.75, y, BREITE - 2.75, h, [_teile(_zitat(e["text"]))], g, fett=True,
                       farbe=self.farben["satzanfang"])

    def _hoehe_feedback(self, e, g):
        halb = BREITE / 2 + 0.5
        return 3.0 + 0.5 + max(_texthoehe([_teile(_zitat(s)) for s in e[seite]], halb, g, abstand_pt=4)
                               for seite in ("gut", "schlecht"))

    def _zeichne_feedback(self, f, e, y, h, g):
        halb = BREITE / 2
        for i, (seite, icon, farbe) in enumerate((("gut", "smiley", self.farben["satzanfang"]),
                                                  ("schlecht", "traurig", self.farben["akzent"]))):
            x = X0 + i * halb
            self._icon(f, icon, x + (halb - 3.0) / 2, y, 3.0, farbe)
            self._textfeld(f, x - 0.5 + i * 0.5, y + 3.5, halb + 0.5, h - 3.5, [_teile(_zitat(s)) for s in e[seite]], g,
                           ausrichtung=PP_ALIGN.CENTER, abstand_pt=4)

    def _hoehe_legende(self, e, g):
        return _cm(g * ZEILE)

    def _zeichne_legende(self, f, e, y, h, g):
        w = BREITE / len(e["eintraege"])
        for i, x in enumerate(e["eintraege"]):
            teil = {"text": x["text"], "fett": True, "farbe": x["farbe"], "unterstrichen": True}
            self._textfeld(f, X0 + 0.21 + i * w, y, w - 0.21, h, [[teil]], g)


def _xml_aufzaehlung(p, einzug):
    """Aufzählungspunkt • mit hängendem Einzug (python-pptx kann das nicht direkt)."""
    ppr = p._p.get_or_add_pPr()
    ppr.set("marL", str(int(einzug)))
    ppr.set("indent", str(-int(einzug)))
    for alt in ppr.findall(qn("a:buNone")):
        ppr.remove(alt)
    zeichen = ppr.makeelement(qn("a:buChar"), {"char": "•"})
    davor = ppr.find(qn("a:tabLst")) if ppr.find(qn("a:tabLst")) is not None else ppr.find(qn("a:defRPr"))
    if davor is not None:
        davor.addprevious(zeichen)
    else:
        ppr.append(zeichen)

