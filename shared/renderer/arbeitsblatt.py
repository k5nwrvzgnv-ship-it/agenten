"""Arbeitsblatt als docx bauen, aus einer geprüften Beschreibung (pruefung.lesen).

python-docx kann keine Rahmen, Nummerierungs-Neustarts oder Wiederholungszeilen. Das steht hier als
kleines XML (Funktionen mit _xml_ im Namen). Seite: A4, Ränder 2 cm, Stile aus der Standardvorlage.
"""

from pathlib import Path

import docx
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

SEITE_CM = (21.0, 29.7)
RAND_CM = 2.0
TEXTBREITE_CM = SEITE_CM[0] - 2 * RAND_CM
LINIENABSTAND_CM = 0.9  # Höhe einer Schreiblinie, passt für Handschrift
SPRACHE = "de-DE"  # Rechtschreibprüfung beim Bearbeiten in Word

# Word verlangt die Reihenfolge der Kindelemente aus dem Schema: neue Elemente vor diesen einfügen.
_NACH_TBLBORDERS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption", "w:tblDescription")
_NACH_PBDR = ("w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct",
              "w:topLinePunct", "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
              "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
              "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl", "w:divId",
              "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def bauen(beschreibung: dict, ziel: Path) -> Path:
    """Baut das Arbeitsblatt nach `ziel` (.docx). `beschreibung` kommt aus pruefung.lesen()."""
    d = docx.Document()
    _seite_einrichten(d, beschreibung["titel"])
    if beschreibung["kopfzeile"]:
        _kopfzeile(d, beschreibung["kopfzeile"])
    d.add_heading(beschreibung["titel"], level=0)
    aufgabe_nr = 0
    for block in beschreibung["bloecke"]:
        typ = block["typ"]
        if typ == "ueberschrift":
            d.add_heading(block["text"], level=block["ebene"])
        elif typ == "absatz":
            d.add_paragraph(block["text"])
        elif typ == "liste":
            _liste(d, block)
        elif typ == "tabelle":
            _tabelle(d, block)
        else:
            aufgabe_nr += 1
            _aufgabe(d, block, aufgabe_nr)
    d.save(ziel)
    return ziel


def _seite_einrichten(d, titel: str):
    s = d.sections[0]
    s.page_width, s.page_height = Cm(SEITE_CM[0]), Cm(SEITE_CM[1])
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(RAND_CM)
    d.core_properties.title = titel
    d.core_properties.author = ""  # sonst "python-docx"
    d.core_properties.last_modified_by = ""
    rpr = d.styles["Normal"].element.get_or_add_rPr()
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.insert_element_before(lang, "w:eastAsianLayout", "w:specVanish", "w:oMath")
    lang.set(qn("w:val"), SPRACHE)


def _kopfzeile(d, felder: list[str]):
    """Felder nebeneinander, je mit unterstrichenem Tab als Linie zum Ausfüllen."""
    kopf = d.sections[0].header
    breite = TEXTBREITE_CM / len(felder)
    t = kopf.add_table(rows=1, cols=len(felder), width=Cm(TEXTBREITE_CM))
    for zelle, feld in zip(t.rows[0].cells, felder, strict=True):
        zelle.width = Cm(breite)
        p = zelle.paragraphs[0]
        p.paragraph_format.tab_stops.add_tab_stop(Cm(breite - 0.4), WD_TAB_ALIGNMENT.RIGHT)
        p.add_run(f"{feld}: ")
        p.add_run("\t").underline = True
    kopf.paragraphs[0]._p.getparent().remove(kopf.paragraphs[0]._p)  # leerer Absatz vor der Tabelle
    kopf.add_paragraph()  # Word verlangt einen Absatz nach einer Tabelle


def _liste(d, block: dict):
    stil = "List Number" if block["nummeriert"] else "List Bullet"
    absaetze = [d.add_paragraph(punkt, style=stil) for punkt in block["punkte"]]
    if block["nummeriert"]:
        _xml_nummerierung_neu(d, stil, absaetze)


def _tabelle(d, block: dict):
    zeilen = ([block["kopf"]] if block["kopf"] else []) + block["zeilen"]
    t = d.add_table(rows=len(zeilen), cols=len(zeilen[0]))
    _xml_rahmen(t)
    for reihe, werte in zip(t.rows, zeilen, strict=True):
        for zelle, wert in zip(reihe.cells, werte, strict=True):
            zelle.paragraphs[0].add_run(wert)
    if block["kopf"]:
        for zelle in t.rows[0].cells:
            for run in zelle.paragraphs[0].runs:
                run.bold = True
        _xml_zeile_wiederholen(t.rows[0])
    _abstand_nach_tabelle(d)


def _aufgabe(d, block: dict, nr: int):
    p = d.add_paragraph()
    p.add_run(f"Aufgabe {nr}: ").bold = True
    p.add_run(block["text"])
    p.paragraph_format.space_before = Pt(12)
    antwort = block["antwort"]
    if antwort is None:
        return
    p.paragraph_format.keep_with_next = True  # Aufgabe nicht von ihrem Antwortplatz trennen
    if antwort["art"] == "linien":
        for i in range(antwort["anzahl"]):
            linie = d.add_paragraph()
            f = linie.paragraph_format
            f.line_spacing = Cm(LINIENABSTAND_CM)  # Länge heißt: genau diese Höhe
            f.space_before = f.space_after = 0
            f.keep_with_next = i < antwort["anzahl"] - 1
            _xml_linie_unten(linie)
    else:
        t = d.add_table(rows=1, cols=1)
        _xml_rahmen(t)
        reihe = t.rows[0]
        reihe.height, reihe.height_rule = Cm(antwort["hoehe_cm"]), WD_ROW_HEIGHT_RULE.EXACTLY
        _abstand_nach_tabelle(d)


def _abstand_nach_tabelle(d):
    """Leerer Absatz nach jeder Tabelle: Zwei Tabellen direkt hintereinander verschmelzen in Word."""
    d.add_paragraph().paragraph_format.space_after = 0


def _xml_rahmen(t):
    """Einfache schwarze Linien um und in der Tabelle, unabhängig von den Stilen der Vorlage."""
    rahmen = OxmlElement("w:tblBorders")
    for seite in ("top", "left", "bottom", "right", "insideH", "insideV"):
        linie = OxmlElement(f"w:{seite}")
        linie.set(qn("w:val"), "single")
        linie.set(qn("w:sz"), "4")  # Achtelpunkte
        linie.set(qn("w:space"), "0")
        linie.set(qn("w:color"), "000000")
        rahmen.append(linie)
    t._tbl.tblPr.insert_element_before(rahmen, *_NACH_TBLBORDERS)


def _xml_zeile_wiederholen(reihe):
    """Kopfzeile der Tabelle auf jeder Seite wiederholen."""
    kopf = OxmlElement("w:tblHeader")
    kopf.set(qn("w:val"), "true")
    reihe._tr.get_or_add_trPr().append(kopf)


def _xml_linie_unten(p):
    """Schreiblinie als unterer Rahmen. 'between' ist nötig, weil Word Absätze mit gleichem Rahmen
    zu einer Gruppe zusammenfasst und den unteren Rahmen sonst nur unter dem letzten zeichnet."""
    rahmen = OxmlElement("w:pBdr")
    for seite in ("bottom", "between"):
        linie = OxmlElement(f"w:{seite}")
        linie.set(qn("w:val"), "single")
        linie.set(qn("w:sz"), "4")
        linie.set(qn("w:space"), "1")
        linie.set(qn("w:color"), "808080")
        rahmen.append(linie)
    p._p.get_or_add_pPr().insert_element_before(rahmen, *_NACH_PBDR)


def _xml_nummerierung_neu(d, stil: str, absaetze):
    """Nummerierte Liste bei 1 beginnen. Ohne das zählt Word über alle Listen des Stils weiter."""
    nummerierung = d.part.numbering_part.element
    stil_num = d.styles[stil].element.pPr.numPr.numId.val
    abstrakt = nummerierung.num_having_numId(stil_num).abstractNumId.val
    neu = nummerierung.add_num(abstrakt)
    neu.add_lvlOverride(ilvl=0).add_startOverride(1)
    for p in absaetze:
        p._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().val = neu.numId
