import tempfile
import unittest
from pathlib import Path

import kurse


class TestKurse(unittest.TestCase):
    def test_sicherer_name(self):
        self.assertEqual(kurse.sicherer_name("../../etc/passwd"), "_.._etc_passwd")
        self.assertEqual(kurse.sicherer_name(".."), "_")
        self.assertEqual(kurse.sicherer_name("Folie 3: Holz/Beton?.pdf"), "Folie 3_ Holz_Beton_.pdf")
        self.assertEqual(kurse.sicherer_name(""), "_")
        self.assertEqual(kurse.sicherer_name("Aufgaben &amp; Einreichungen"), "Aufgaben & Einreichungen")
        self.assertEqual(kurse.sicherer_name("&#47;etc"), "_etc")

    def test_html_zu_text(self):
        self.assertEqual(kurse.html_zu_text("<p>Klausur am <b>12.02.</b></p><p>Raum&nbsp;5</p>"),
                         "Klausur am 12.02.\nRaum\xa05")
        self.assertEqual(kurse.html_zu_text(None), "")

    def test_textfassungen(self):
        import docx
        import openpyxl
        import pptx
        ordner = Path(tempfile.mkdtemp())
        d = docx.Document()
        d.add_paragraph("Hausarbeit Holzschutz")
        d.save(ordner / "a.docx")
        p = pptx.Presentation()
        for titel in ("Einleitung", "Brettsperrholz"):
            p.slides.add_slide(p.slide_layouts[0]).shapes.title.text = titel
        p.save(ordner / "b.pptx")
        w = openpyxl.Workbook()
        w.active.title = "Messwerte"
        w.active.append(["Feuchte", 12.5])
        w.save(ordner / "c.xlsx")
        self.assertEqual(kurse.als_text(ordner / "a.docx"), "Hausarbeit Holzschutz")
        self.assertEqual(kurse.als_text(ordner / "b.pptx"), "--- Folie 1\nEinleitung\n--- Folie 2\nBrettsperrholz")
        self.assertEqual(kurse.als_text(ordner / "c.xlsx"), "--- Blatt Messwerte\nFeuchte | 12.5")
        self.assertTrue(kurse.textfassung(ordner / "b.pptx", erneuern=False))
        self.assertTrue((ordner / "b.pptx.txt").exists())
        (ordner / "x.png").write_bytes(b"png")
        self.assertFalse(kurse.textfassung(ordner / "x.png", erneuern=False))

    def test_aufgabe_als_text(self):
        a = {"name": "GEG", "cmid": 5, "duedate": 0, "cutoffdate": 0, "allowsubmissionsfromdate": 0,
             "teamsubmission": 1, "intro": "<p>Hüllfläche finalisieren</p>",
             "configs": [{"plugin": "onlinetext", "subtype": "assignsubmission", "name": "enabled", "value": "1"},
                         {"plugin": "file", "subtype": "assignsubmission", "name": "enabled", "value": "1"},
                         {"plugin": "file", "subtype": "assignsubmission", "name": "maxfilesubmissions", "value": "2"},
                         {"plugin": "comments", "subtype": "assignfeedback", "name": "enabled", "value": "1"}]}
        text = "\n".join(kurse.aufgabe_als_text(a, "nichts abgegeben", []))
        self.assertIn("Abgabeart: Online-Text (direkt in Moodle eintippen); Datei-Upload, höchstens 2 Datei(en)", text)
        self.assertIn("Teamabgabe: ja", text)
        self.assertIn("Hüllfläche finalisieren", text)
        self.assertNotIn("Jans Gruppe", text)  # Gruppen unbekannt: keine Zeile
        mit = "\n".join(kurse.aufgabe_als_text(a, "nichts abgegeben", [], gruppen=[]))
        self.assertIn("Jans Gruppe im Kurs: KEINE", mit)
        mit = "\n".join(kurse.aufgabe_als_text(a, "nichts abgegeben", [], gruppen=["Gr.2 gelb"]))
        self.assertIn("Jans Gruppe im Kurs: Gr.2 gelb", mit)
        a["teamsubmission"] = 0
        self.assertNotIn("Jans Gruppe", "\n".join(kurse.aufgabe_als_text(a, "x", [], gruppen=[])))

    def test_gruppenwahl_als_text(self):
        zeilen = kurse.gruppenwahl_als_text([
            {"name": "Gr.2 gelb", "maxanswers": 7, "countanswers": 0, "checked": False},
            {"name": "Gr. 7  weiss (Voll belegt) ", "maxanswers": 6, "countanswers": 6, "checked": True},
            {"name": "offen", "maxanswers": 0, "countanswers": 3, "checked": False}])
        self.assertEqual(zeilen[1:], ["  - Gr.2 gelb: 0 von 7, 7 frei",
                                      "  - Gr. 7  weiss (Voll belegt): 6 von 6, 0 frei <- Jans Wahl",
                                      "  - offen: 3 (ohne Grenze)"])
        self.assertEqual(kurse.gruppenwahl_als_text([]), ["  Belegung: (noch keine Gruppen angelegt)"])

    def test_pdf_ohne_text(self):
        import pypdf
        ordner = Path(tempfile.mkdtemp())
        w = pypdf.PdfWriter()
        w.add_blank_page(100, 100)
        w.write(ordner / "scan.pdf")
        self.assertIn("vermutlich ein Scan", kurse.als_text(ordner / "scan.pdf"))

    def test_entpacken_bleibt_im_ordner(self):
        import zipfile
        ordner = Path(tempfile.mkdtemp())
        with zipfile.ZipFile(ordner / "a.zip", "w") as z:
            z.writestr("../../boese.txt", "x")
            z.writestr("unter/ordner/daten.csv", "1;2")
            z.writestr("/absolut.txt", "y")
        self.assertTrue(kurse.textfassung(ordner / "a.zip", erneuern=True))
        dateien = sorted(str(p.relative_to(ordner)) for p in ordner.rglob("*") if p.is_file())
        self.assertEqual(dateien, ["a.zip", "a.zip.txt", "a_entpackt/_/_/boese.txt", "a_entpackt/absolut.txt",
                                   "a_entpackt/unter/ordner/daten.csv"])
        for p in ordner.rglob("*"):
            self.assertTrue(p.resolve().is_relative_to(ordner.resolve()))
        self.assertTrue((ordner / "a_entpackt" / "unter" / "ordner" / "daten.csv").exists())
        self.assertIn("unter/ordner/daten.csv", (ordner / "a.zip.txt").read_text())


if __name__ == "__main__":
    unittest.main()


class FakeMoodle:
    """Zwei Kurse mit je einer Datei, ohne Netz."""
    url = "https://moodle.test"

    def token(self):
        return "t"

    def kurse(self):
        return [{"id": 1, "fullname": "Alt"}, {"id": 2, "fullname": "Neu"}]

    def aufruf(self, funktion, **p):
        if funktion == "mod_assign_get_assignments":
            return {"courses": []}
        datei = {"type": "file", "filename": f"k{p['courseid']}.pdf", "fileurl": f"u{p['courseid']}",
                 "timemodified": 2, "filesize": 10}
        return [{"name": "A", "modules": [{"id": 9, "modname": "resource", "name": "R", "contents": [datei]}]}]


class TestSpiegeln(unittest.TestCase):
    def test_neuer_kurs_meldet_nichts(self):
        with tempfile.TemporaryDirectory() as tmp:
            alt = (kurse.ZIEL, kurse.STAND_DATEI, kurse._laden, kurse.textfassung)
            kurse.ZIEL = Path(tmp)
            kurse.STAND_DATEI = Path(tmp) / ".stand.json"
            kurse._laden = lambda url, token, ziel: (ziel.parent.mkdir(parents=True, exist_ok=True),
                                                     ziel.write_bytes(b"x"))
            kurse.textfassung = lambda pfad, erneuern: False
            try:
                kurse.STAND_DATEI.write_text("{}")
                (Path(tmp) / "Alt").mkdir()
                (Path(tmp) / "Alt" / "uebersicht.md").write_text("")
                neu = kurse.spiegeln(FakeMoodle())
                self.assertEqual(neu, ["Alt/00 A/k1.pdf"])
                self.assertTrue((Path(tmp) / "Neu" / "00 A" / "k2.pdf").exists())
                self.assertEqual(kurse.spiegeln(FakeMoodle()), [])
            finally:
                kurse.ZIEL, kurse.STAND_DATEI, kurse._laden, kurse.textfassung = alt
