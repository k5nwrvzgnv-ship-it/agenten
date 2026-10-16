"""Kursmaterial aus Moodle nach data/kurse/ spiegeln, damit Claude es lesen kann.

Pro Kurs: data/kurse/<Kurs>/uebersicht.md (Abschnitte, Texte, Links, Dateipfade)
und die Dateien in Unterordnern je Abschnitt. Nur neue oder geänderte Dateien
werden geladen (data/kurse/.stand.json: Datei-URL -> timemodified).
Nicht geladen: ZIP-Archive und Dateien über MAX_MB (stehen nur in der Übersicht).
Word, PowerPoint und Excel (docx, pptx, xlsx) bekommen daneben eine Textfassung
<datei>.txt, weil Claude sie sonst nicht lesen kann. Alte Formate (doc, ppt, xls) nicht.
In Moodle gelöschte Dateien bleiben liegen.
"""

import json
import re
import shutil
import urllib.request
from html import unescape
from pathlib import Path

HIER = Path(__file__).resolve().parent
ZIEL = HIER / "data" / "kurse"
STAND_DATEI = ZIEL / ".stand.json"
MAX_MB = 100
NICHT_LADEN = (".zip",)


def sicherer_name(name: str) -> str:
    """Ein Pfadteil ohne Schrägstriche, Steuerzeichen und führende Punkte."""
    name = re.sub(r'[\x00-\x1f/\\:*?"<>|]', "_", name).strip().lstrip(".")
    return name[:120] or "_"


def html_zu_text(html: str) -> str:
    text = re.sub(r"<br\s*/?>|</p>|</li>|</h\d>", "\n", html or "", flags=re.I)
    text = unescape(re.sub(r"<[^>]+>", "", text))
    return "\n".join(z.strip() for z in text.splitlines() if z.strip())


def als_text(pfad: Path) -> str | None:
    """Text aus docx, pptx oder xlsx (mit Folien- bzw. Blattangabe). Andere Formate: None."""
    endung = pfad.suffix.lower()
    if endung == ".docx":
        import docx
        d = docx.Document(pfad)
        teile = [a.text for a in d.paragraphs if a.text.strip()]
        for t in d.tables:
            teile += [" | ".join(z.text.strip() for z in reihe.cells) for reihe in t.rows]
        return "\n".join(teile)
    if endung == ".pptx":
        import pptx
        teile = []
        for nr, folie in enumerate(pptx.Presentation(pfad).slides, 1):
            teile.append(f"--- Folie {nr}")
            for form in folie.shapes:
                if form.has_text_frame and form.text_frame.text.strip():
                    teile.append(form.text_frame.text.strip())
                if getattr(form, "has_table", False) and form.has_table:
                    teile += [" | ".join(z.text.strip() for z in reihe.cells) for reihe in form.table.rows]
            if folie.has_notes_slide and folie.notes_slide.notes_text_frame.text.strip():
                teile.append(f"Notizen: {folie.notes_slide.notes_text_frame.text.strip()}")
        return "\n".join(teile)
    if endung == ".xlsx":
        import openpyxl
        teile = []
        for blatt in openpyxl.load_workbook(pfad, read_only=True, data_only=True).worksheets:
            teile.append(f"--- Blatt {blatt.title}")
            for reihe in blatt.iter_rows(values_only=True):
                if any(w is not None for w in reihe):
                    teile.append(" | ".join("" if w is None else str(w) for w in reihe))
        return "\n".join(teile)
    return None


def _textfassung(pfad: Path, erneuern: bool) -> bool:
    """Schreibt <pfad>.txt, falls das Format unterstützt wird. True, wenn es eine Textfassung gibt."""
    if pfad.suffix.lower() not in (".docx", ".pptx", ".xlsx"):
        return False
    txt = pfad.with_name(pfad.name + ".txt")
    if erneuern or not txt.exists():
        try:
            txt.write_text(als_text(pfad) or "", encoding="utf-8")
        except Exception as f:  # noqa: BLE001 – kaputte Datei soll den Abgleich nicht stoppen
            txt.write_text(f"(Textfassung fehlgeschlagen: {f})", encoding="utf-8")
    return True


def _laden(url: str, token: str, ziel: Path) -> None:
    trenner = "&" if "?" in url else "?"
    ziel.parent.mkdir(parents=True, exist_ok=True)
    tmp = ziel.with_name(f".{ziel.name}.tmp")
    with urllib.request.urlopen(f"{url}{trenner}token={token}", timeout=300) as antwort, open(tmp, "wb") as f:
        shutil.copyfileobj(antwort, f)
    tmp.replace(ziel)


def spiegeln(m) -> list[str]:
    """Gleicht alle Kurse ab. Gibt neue oder geänderte Dateien zurück (relativ zu ZIEL).

    Beim ersten Abgleich (noch kein Stand) kommt eine leere Liste zurück, damit
    nicht der ganze Bestand als "neu" gemeldet wird.
    """
    erster = not STAND_DATEI.exists()
    stand = json.loads(STAND_DATEI.read_text()) if not erster else {}
    neu = []
    vergeben = set()
    kurse = m.kurse()
    namen = [k["fullname"] for k in kurse]
    for k in kurse:
        kursname = k["fullname"] if namen.count(k["fullname"]) == 1 else f"{k['fullname']} ({k['id']})"
        kursordner = ZIEL / sicherer_name(kursname)
        zeilen = [f"# {k['fullname']}", "", f"Moodle: {m.url}/course/view.php?id={k['id']}", ""]
        for nr, abschnitt in enumerate(m.aufruf("core_course_get_contents", courseid=k["id"])):
            titel = abschnitt.get("name") or f"Abschnitt {nr}"
            abschnitt_ordner = f"{nr:02d} {sicherer_name(titel)}"
            zeilen += [f"## {titel}", ""]
            if text := html_zu_text(abschnitt.get("summary", "")):
                zeilen += [text, ""]
            for mod in abschnitt["modules"]:
                zeilen.append(f"- {mod['modname']}: {mod.get('name', '')}")
                if mod["modname"] == "url":
                    zeilen += [f"  Link: {c['fileurl']}" for c in mod.get("contents") or [] if c.get("fileurl")]
                if text := html_zu_text(mod.get("description", "")):
                    zeilen += [f"  {z}" for z in text.splitlines()]
                for c in mod.get("contents") or []:
                    if c.get("type") != "file":
                        continue
                    teile = [sicherer_name(t) for t in (c.get("filepath") or "/").split("/") if t]
                    unterordner = [sicherer_name(mod["name"])] if mod["modname"] == "folder" else []
                    rel = Path(sicherer_name(kursname), abschnitt_ordner, *unterordner, *teile,
                               sicherer_name(c["filename"]))
                    if rel in vergeben:  # gleicher Dateiname in zwei Moodle-Einträgen eines Abschnitts
                        rel = rel.with_name(f"{rel.stem} ({mod['id']}){rel.suffix}")
                    vergeben.add(rel)
                    mb = c.get("filesize", 0) / 1e6
                    if c["filename"].lower().endswith(NICHT_LADEN) or mb > MAX_MB:
                        zeilen.append(f"  Datei (nicht geladen, {mb:.0f} MB): {c['filename']}")
                        continue
                    zeilen.append(f"  Datei: {rel.relative_to(sicherer_name(kursname))}")
                    schluessel = c["fileurl"]
                    geladen = stand.get(schluessel) != c.get("timemodified") or not (ZIEL / rel).exists()
                    if geladen:
                        _laden(c["fileurl"], m.token(), ZIEL / rel)
                        stand[schluessel] = c.get("timemodified")
                        neu.append(str(rel))
                        _speichern(stand)
                    if _textfassung(ZIEL / rel, erneuern=geladen):
                        zeilen.append(f"  Textfassung: {rel.relative_to(sicherer_name(kursname))}.txt")
            zeilen.append("")
        kursordner.mkdir(parents=True, exist_ok=True)
        (kursordner / "uebersicht.md").write_text("\n".join(zeilen), encoding="utf-8")
    _speichern(stand)
    return [] if erster else neu


def _speichern(stand: dict) -> None:
    ZIEL.mkdir(parents=True, exist_ok=True)
    tmp = STAND_DATEI.with_suffix(".tmp")
    tmp.write_text(json.dumps(stand))
    tmp.replace(STAND_DATEI)
