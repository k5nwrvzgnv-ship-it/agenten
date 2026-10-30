import os
import tempfile
import unittest
import zipfile
from pathlib import Path

import docx

import formatvorlage


def als_dotx(quelle: Path, ziel: Path) -> None:
    """docx -> dotx (Inhaltstyp der Vorlage), wie Word sie speichert."""
    with zipfile.ZipFile(quelle) as rein, zipfile.ZipFile(ziel, "w") as raus:
        for e in rein.infolist():
            daten = rein.read(e.filename)
            if e.filename == "[Content_Types].xml":
                daten = daten.replace(b"wordprocessingml.document.main+xml", b"wordprocessingml.template.main+xml")
            raus.writestr(e, daten)


class FormatvorlageTest(unittest.TestCase):
    def setUp(self):
        self.hier = Path(tempfile.mkdtemp())
        self.ordner = self.hier / "data" / "vorlagen"

    def test_ohne_ordner_oder_vorlage(self):
        self.assertEqual(formatvorlage.waehlen(self.hier, ".docx"), (None, []))
        self.ordner.mkdir(parents=True)
        (self.ordner / "notiz.txt").write_text("x")
        self.assertEqual(formatvorlage.waehlen(self.hier, ".docx"), (None, []))
        self.assertEqual(formatvorlage.waehlen(self.hier, ".xlsx"), (None, []))

    def test_dotx_wird_umgewandelt_neueste_gilt(self):
        self.ordner.mkdir(parents=True)
        docx.Document().save(self.ordner / "alt.docx")
        os.utime(self.ordner / "alt.docx", (1, 1))
        als_dotx(self.ordner / "alt.docx", self.ordner / "Schule.dotx")
        (self.ordner / "~$Schule.dotx").write_bytes(b"Sperrdatei von Word")
        vorlage, hinweise = formatvorlage.waehlen(self.hier, ".docx")
        self.assertEqual(vorlage, self.ordner / "umgewandelt" / "Schule.docx")
        docx.Document(str(vorlage))  # öffnet
        self.assertEqual(hinweise, ["Vorlage Schule.dotx", "weitere Vorlagen nicht genutzt: alt.docx"])
        self.assertEqual(formatvorlage.waehlen(self.hier, ".pptx"), (None, []))

    def test_vorlage_je_art(self):
        (self.ordner / "reihenentwurf").mkdir(parents=True)
        docx.Document().save(self.ordner / "allgemein.docx")
        self.assertEqual(formatvorlage.waehlen(self.hier, ".docx", "reihenentwurf")[0], self.ordner / "allgemein.docx")
        docx.Document().save(self.ordner / "reihenentwurf" / "entwurf.docx")
        vorlage, hinweise = formatvorlage.waehlen(self.hier, ".docx", "reihenentwurf")
        self.assertEqual(vorlage, self.ordner / "reihenentwurf" / "entwurf.docx")
        self.assertEqual(formatvorlage.waehlen(self.hier, ".docx")[0], self.ordner / "allgemein.docx")
        self.assertEqual(formatvorlage.waehlen(self.hier, ".docx", "lernerfolgskontrolle")[0],
                         self.ordner / "allgemein.docx")

    def test_kaputte_vorlage(self):
        self.ordner.mkdir(parents=True)
        (self.ordner / "kaputt.dotx").write_bytes(b"kein zip")
        with self.assertRaises(ValueError) as f:
            formatvorlage.waehlen(self.hier, ".docx")
        self.assertIn("kaputt.dotx", str(f.exception))
        self.assertFalse((self.ordner / "umgewandelt" / "kaputt.docx").exists())


if __name__ == "__main__":
    unittest.main()
